"""정답지 엑셀을 채점기가 읽는 정답표 JSON으로 바꾸고, 정답지가 바뀌었는지 감시한다(묶음 ⑯, 기준은 CHANGES.md의 N1~N12).

흐름:
  convert(엑셀, 출력 폴더)   모든 검사를 통과해야만 JSON을 쓴다. 처음이면 스냅샷은 "미확인"이다.
  confirm(출력 폴더, 사람, 사유)  사람이 확인했다고 기록한다("확인됨").
  load_confirmed(출력 폴더)  판정과 실행기가 정답표를 읽는 유일한 입구. "확인됨"이고 JSON이 그대로일 때만 돌려준다.
확인된 뒤 정답지가 달라지면 convert는 무엇이 달라졌는지 알리고 멈춘다. 사람이 사유와 함께 받아들일 때만 새로 쓴다.
원본 엑셀은 읽기만 한다.
"""
from __future__ import annotations

import datetime
import hashlib
import json
import os
import re
import sys
import unicodedata

import scoring

HEADER_ROW, FIRST_ROW = 4, 5
SHEET_Q, SHEET_E = "질문 정답표", "질문별 근거"
COL_Q = {"질문 ID": "qid", "유형": "type", "질문": "text", "답변 가능": "ans", "필수 근거 수": "n_req", "보조 근거 수": "n_opt"}
COL_E = {"질문 ID": "qid", "파일": "file", "시트": "sheet", "시작 행": "start", "끝 행": "end", "필수/보조": "role"}
JSON_NAME, SNAP_NAME = "gold_questions.json", "gold_snapshot.json"
UNCONFIRMED, CONFIRMED = "미확인", "확인됨"


class XlsxGoldError(ValueError):
    """정답지 엑셀이 명세에 맞지 않을 때."""


class GoldChanged(XlsxGoldError):
    """확인된 정답지와 달라져서 사람의 승인을 기다릴 때."""


class GoldNotConfirmed(XlsxGoldError):
    """확인되지 않았거나 확인 뒤에 JSON이 바뀌어 판정에 쓸 수 없을 때."""


# ---------- 읽기와 검사 (N1~N3) ----------
def _table(ws, cols, label, optional=None):
    """필수 열(cols)은 없으면 멈추고, 선택 열(optional)은 머리글에 있을 때만 읽는다(묶음 ⑲ Q1). 읽은 선택 열의 키 집합도 돌려준다."""
    names = [c.value.strip() if isinstance(c.value, str) else c.value for c in ws[HEADER_ROW]]
    seen = [n for n in names if n]
    dup = sorted({n for n in seen if seen.count(n) > 1})
    if dup:
        raise XlsxGoldError(f"{label}: 머리글({HEADER_ROW}행)에 같은 열 이름이 두 번 이상 있습니다 {dup}")
    header = {n: i for i, n in enumerate(names) if n}
    missing = [n for n in cols if n not in header]
    if missing:
        raise XlsxGoldError(f"{label}: 머리글({HEADER_ROW}행)에 없는 열 이름 {missing}")
    present = {name: key for name, key in (optional or {}).items() if name in header}
    rows = []
    for r in ws.iter_rows(min_row=FIRST_ROW):
        cells = [c.value for c in r]
        if cells[header["질문 ID"]] in (None, ""):
            if any(v not in (None, "") for v in cells):               # 질문 ID만 비고 다른 칸에 내용이 있으면 조용히 넘기지 않는다
                raise XlsxGoldError(f"{label} {r[0].row}행: 질문 ID가 비었는데 다른 칸에 내용이 있습니다")
            continue
        rows.append((r[0].row, {key: cells[header[name]] for name, key in {**cols, **present}.items()}))
    unused = sorted(n for n in header if n not in cols and n not in present)
    return rows, unused, set(present.values())


def _text(v, where, what, required=False):
    """선택 열의 글 칸. 비면 None(required면 멈춤), 글이 아니면 멈춘다. 앞뒤 공백은 지운다."""
    if v is None or (isinstance(v, str) and not v.strip()):
        if required:
            raise XlsxGoldError(f"{where}: {what} 칸이 비었습니다")
        return None
    if not isinstance(v, str):
        raise XlsxGoldError(f"{where}: {what}는 글이어야 합니다: {v!r}")
    return v.strip()


def _whole(v, where, what):
    if isinstance(v, bool) or not isinstance(v, (int, float)) or float(v) != int(v):
        raise XlsxGoldError(f"{where}: {what}는 정수여야 합니다: {v!r}")
    return int(v)


def _file_name(v, where):
    if not isinstance(v, str) or not v.strip():
        raise XlsxGoldError(f"{where}: 파일 칸이 비었습니다")
    if "/" in v or "\\" in v:
        raise XlsxGoldError(f"{where}: 파일은 경로 없이 이름만 써야 합니다: {v!r}")
    if not unicodedata.is_normalized("NFC", v):
        raise XlsxGoldError(f"{where}: 파일 이름이 분해형 한글(NFD)입니다. 실제 파일 이름을 복사해 다시 쓰세요: {v!r}")
    return v


def read_gold_xlsx(path: str):
    """(정답표 목록, 변환 보고)를 돌려준다. 어긋나면 XlsxGoldError."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)            # 수식 칸은 저장된 값을 읽는다
    for s in (SHEET_Q, SHEET_E):
        if s not in wb.sheetnames:
            raise XlsxGoldError(f"시트가 없습니다: {s}")
    fmt = scoring.load_format()                                   # 묶음 ⑲ Q18: 유형 목록과 선택 열 이름
    fc = fmt["columns"]
    qrows, unused_q, got_q = _table(wb[SHEET_Q], COL_Q, SHEET_Q, {fc["question_conditions"]: "stated"})
    erows, unused_e, got_e = _table(wb[SHEET_E], COL_E, SHEET_E, {fc["evidence_condition"]: "condition", fc["evidence_area"]: "area",
                                                                   fc["evidence_bundle"]: "bundle", fc["evidence_requirement"]: "requirement"})
    if ("stated" in got_q) != ("condition" in got_e):                # Q1: 조건 열은 짝으로 쓴다
        raise XlsxGoldError(f"'{fc['question_conditions']}'({SHEET_Q})과 '{fc['evidence_condition']}'({SHEET_E})은 짝으로 있어야 합니다(한쪽만 있음)")
    questions, by_id = [], {}
    for rn, q in qrows:
        qid, where = q["qid"], f"{SHEET_Q} {rn}행({q['qid']})"
        if qid in by_id:
            raise XlsxGoldError(f"{where}: 질문 ID 중복")
        if q["type"] not in fmt["types"]:
            raise XlsxGoldError(f"{where}: 유형이 {tuple(fmt['types'])} 중 하나가 아닙니다: {q['type']!r}")
        if q["ans"] not in ("O", "X"):
            raise XlsxGoldError(f"{where}: 답변 가능은 O 또는 X여야 합니다: {q['ans']!r}")
        answerable = q["ans"] == "O"
        if (q["type"] == fmt["unanswerable_type"]) == answerable:
            raise XlsxGoldError(f"{where}: 유형({q['type']})과 답변 가능({q['ans']})이 서로 맞지 않습니다")
        for k, label in (("n_req", "필수 근거 수"), ("n_opt", "보조 근거 수")):
            if q[k] is None:
                raise XlsxGoldError(f"{where}: {label} 칸에 저장된 값이 없습니다(수식이면 엑셀에서 열어 저장한 파일을 쓰세요)")
            q[k] = _whole(q[k], where, label)
        d = {"case_id": qid, "type": q["type"], "request": q["text"], "answerable": answerable, "gold_evidence": []}
        if "stated" in got_q:                                       # Q12: 새 칸은 그 열이 있을 때만 쓴다(옛 정답표의 JSON은 그대로)
            s = _text(q["stated"], where, fc["question_conditions"])
            d["stated_conditions"] = [x.strip() for x in s.split(fmt["condition_separator"]) if x.strip()] if s else []
        by_id[qid] = (d, q)
        questions.append(d)
    counts = {qid: {"필수": 0, "보조": 0} for qid in by_id}
    for rn, e in erows:
        where = f"{SHEET_E} {rn}행({e['qid']})"
        if e["qid"] not in by_id:
            raise XlsxGoldError(f"{where}: 질문 정답표에 없는 질문 ID")
        if e["role"] not in ("필수", "보조"):
            raise XlsxGoldError(f"{where}: 필수/보조 칸이 잘못되었습니다: {e['role']!r}")
        if not isinstance(e["sheet"], str) or not e["sheet"].strip():
            raise XlsxGoldError(f"{where}: 시트 칸이 비었습니다")
        s, t = _whole(e["start"], where, "시작 행"), _whole(e["end"], where, "끝 행")
        if s < 1 or t < s:
            raise XlsxGoldError(f"{where}: 행 범위는 1 이상이고 시작 <= 끝이어야 합니다: {s}~{t}")
        ev = {"file": _file_name(e["file"], where), "sheet": e["sheet"], "rows": [s, t], "required": e["role"] == "필수"}
        if "condition" in got_e:                                    # Q2: 값을 만들어 채우지 않는다. 비면 멈춘다(조건과 무관하면 '항상')
            ev["condition"] = _text(e["condition"], where, fc["evidence_condition"], required=True)
        if "area" in got_e:
            ev["area"] = _text(e["area"], where, fc["evidence_area"])
        if "bundle" in got_e:
            ev["bundle"] = _text(e["bundle"], where, fc["evidence_bundle"])
        if "requirement" in got_e:                                  # Q16: O/X가 아니면 멈춘다
            if e["requirement"] not in ("O", "X"):
                raise XlsxGoldError(f"{where}: {fc['evidence_requirement']}는 O 또는 X여야 합니다: {e['requirement']!r}")
            ev["requirement"] = e["requirement"] == "O"
        by_id[e["qid"]][0]["gold_evidence"].append(ev)
        counts[e["qid"]][e["role"]] += 1
    for qid, (d, q) in by_id.items():
        got, want = (counts[qid]["필수"], counts[qid]["보조"]), (q["n_req"], q["n_opt"])
        if got != want:
            raise XlsxGoldError(f"질문 {qid}: 근거 시트의 필수·보조 수 {got[0]}·{got[1]}가 질문 정답표의 {want[0]}·{want[1]}과 다릅니다")
    if not questions:
        raise XlsxGoldError(f"{SHEET_Q}에 질문이 하나도 없습니다")
    try:
        scoring.parse_gold(questions)                           # 채점기의 규칙(필수 근거 1~8개 등)을 그대로 통과해야 한다
    except scoring.ScoringError as err:
        raise XlsxGoldError(f"채점기 검사에서 멈춤: {err}") from None
    return questions, {"unused_columns": {SHEET_Q: unused_q, SHEET_E: unused_e}}


# ---------- 묶음 ㉒: 평가셋 v0.5 읽기(--format, 기준은 CHANGES.md의 V1~V26) ----------
_PROFILE_KEYS = {"profile", "rule", "sheets", "header_row", "first_row", "question_columns", "row_kind_read", "sub_type_none", "answerable_values",
                 "types", "evidence_columns", "role_values", "condition_expected", "document_columns", "documents"}
_QCOLS = ("id", "row_kind", "main_type", "sub_type", "text", "answerable", "auto_required", "auto_conditional", "auto_optional", "state_review", "state_kind")
_ECOLS = ("question_id", "unit", "bundle", "role", "document", "sheet", "cells", "area", "condition")
_DCOLS = ("id", "sha256", "state")


def load_profile(path: str) -> dict:
    """읽기 설정을 읽는다. 칸이 하나라도 빠지거나 모르는 칸이 있으면 멈춘다(조용한 실패 금지)."""
    import yaml

    with open(path, encoding="utf-8") as f:
        raw = yaml.safe_load(f)
    if not isinstance(raw, dict):
        raise XlsxGoldError(f"{path}: 설정은 키와 값의 목록이어야 합니다")
    extra, missing = set(raw) - _PROFILE_KEYS, _PROFILE_KEYS - set(raw)
    if extra or missing:
        raise XlsxGoldError(f"{path}: 칸이 맞지 않습니다(모르는 칸 {sorted(extra)}, 빠진 칸 {sorted(missing)})")
    if raw["rule"] != scoring.UNIT_BUNDLE:
        raise XlsxGoldError(f"{path}: rule은 {scoring.UNIT_BUNDLE!r}만 쓸 수 있습니다: {raw['rule']!r}")
    for key, names in (("question_columns", _QCOLS), ("evidence_columns", _ECOLS), ("document_columns", _DCOLS)):
        got = raw[key]
        if not isinstance(got, dict) or set(got) != set(names) or not all(isinstance(v, str) and v.strip() for v in got.values()):
            raise XlsxGoldError(f"{path}: {key}는 {list(names)}를 모두 가진 비어 있지 않은 문자열이어야 합니다")
    if not isinstance(raw["documents"], dict) or not raw["documents"]:
        raise XlsxGoldError(f"{path}: documents가 비어 있습니다")
    for doc, v in raw["documents"].items():
        if not isinstance(v, dict) or not {"file", "sha256"} <= set(v) or set(v) - {"file", "sha256", "release"}                 or not all(isinstance(v[k], str) and v[k].strip() for k in ("file", "sha256")):
            raise XlsxGoldError(f"{path}: documents.{doc}는 file과 sha256(선택으로 release)을 가져야 합니다")
        _file_name(v["file"], f"{path} documents.{doc}")
        rel = v.get("release")                                      # 묶음 ㉒ V27: 보류를 사람이 푸는 기록(by, date, reason 모두 필요)
        if "release" in v and (not isinstance(rel, dict) or set(rel) != {"by", "date", "reason"}
                               or not all(isinstance(x, str) and x.strip() for x in rel.values())):
            raise XlsxGoldError(f"{path}: documents.{doc}.release는 by, date, reason을 모두 가진 비어 있지 않은 글이어야 합니다")
    for key in ("answerable_values", "role_values"):
        if not isinstance(raw[key], dict) or not raw[key] or not all(isinstance(b, bool) for b in raw[key].values()):
            raise XlsxGoldError(f"{path}: {key}는 값 이름 -> true/false 목록이어야 합니다")
    if not isinstance(raw["types"], list) or not raw["types"]:
        raise XlsxGoldError(f"{path}: types가 비어 있습니다")
    if not all(isinstance(raw["sheets"].get(k), str) for k in ("questions", "evidence", "documents")):
        raise XlsxGoldError(f"{path}: sheets는 questions, evidence, documents 이름을 가져야 합니다")
    for k in ("header_row", "first_row"):
        if isinstance(raw[k], bool) or not isinstance(raw[k], int) or raw[k] < 1:
            raise XlsxGoldError(f"{path}: {k}는 1 이상의 정수여야 합니다")
    return raw


def _profile_header(ws, header_row: int, wanted: dict, label: str) -> dict:
    """머리글에서 설정의 열 이름을 모두 찾는다. 없거나 같은 이름이 둘 이상이면 멈추고 이름을 적는다(V20). 비슷한 이름으로 추측하지 않는다."""
    names = [c.value.strip() if isinstance(c.value, str) else c.value for c in ws[header_row]]
    seen = [n for n in names if n]
    dup = sorted({n for n in seen if seen.count(n) > 1})
    missing = sorted(v for v in wanted.values() if v not in names)
    if missing or any(v in dup for v in wanted.values()):
        raise XlsxGoldError(f"{label}: 머리글({header_row}행)에서 설정한 열 이름을 못 찾거나 같은 이름이 둘 이상입니다: 없음 {missing}, 중복 "
                            f"{sorted(v for v in wanted.values() if v in dup)}")
    return {k: names.index(v) for k, v in wanted.items()}, [n for n in seen if n not in wanted.values()]


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _cell_list(v, where):
    if not isinstance(v, str) or not v.strip():
        raise XlsxGoldError(f"{where}: 셀 범위 칸이 비었거나 글이 아닙니다: {v!r}")
    toks = [t.strip() for t in v.split(",")]
    if any(not t for t in toks):
        raise XlsxGoldError(f"{where}: 셀 범위에 빈 항목이 있습니다: {v!r}")
    try:
        scoring.cells_to_rows(toks, where)                           # 모르는 모양이면 채점기의 검사가 멈춘다
    except scoring.ScoringError as err:
        raise XlsxGoldError(str(err)) from None
    return toks


def read_gold_profile(path: str, profile_path: str):
    """(정답표 목록, 변환 보고)를 돌려준다. 평가셋 v0.5처럼 양식이 다른 정답지를 설정 파일대로 읽는다. 어긋나면 XlsxGoldError.
    원본 문서(마스킹본)는 열지 않는다. 지문은 설정에 적은 값과 평가셋 시트의 값만 비교한다(V25)."""
    import openpyxl

    prof = load_profile(profile_path)
    fmt = scoring.load_format()
    wb = openpyxl.load_workbook(path, data_only=True)
    sh = prof["sheets"]
    for s in sh.values():
        if s not in wb.sheetnames:
            raise XlsxGoldError(f"시트가 없습니다: {s}")
    hdr, first = prof["header_row"], prof["first_row"]
    qi, q_unused = _profile_header(wb[sh["questions"]], hdr, prof["question_columns"], sh["questions"])
    ei, e_unused = _profile_header(wb[sh["evidence"]], hdr, prof["evidence_columns"], sh["evidence"])
    di, d_unused = _profile_header(wb[sh["documents"]], hdr, prof["document_columns"], sh["documents"])

    # 기준 문서 시트: 지문을 비교해 상태를 계산한다(V5, status를 사람이 적지 않는다)
    sheet_docs = {}
    for r in wb[sh["documents"]].iter_rows(min_row=first):
        cells = [c.value for c in r]
        did = cells[di["id"]]
        if _blank(did):
            continue
        if str(did).strip() in sheet_docs:                                        # 묶음 ㉝ Y1: 행 순서로 결과가 바뀌지 않게 중복은 멈춘다
            raise XlsxGoldError(f"{sh['documents']} {r[0].row}행: 문서 ID({str(did).strip()})가 중복입니다")
        sheet_docs[str(did).strip()] = {"sha256": None if _blank(cells[di["sha256"]]) else str(cells[di["sha256"]]).strip().lower(),
                                         "state": None if _blank(cells[di["state"]]) else str(cells[di["state"]]).strip()}
    doc_status = {}
    for did, spec in prof["documents"].items():
        got = sheet_docs.get(did)
        if got is None:
            doc_status[did] = {"status": "unverifiable", "state": None, "reason": "평가셋 기준 문서 시트에 없음"}
        elif got["sha256"] == spec["sha256"].lower():
            doc_status[did] = {"status": "ok", "state": got["state"], "reason": None, "hash_mismatch": False}
            if spec.get("release"):
                doc_status[did]["released_unneeded"] = dict(spec["release"])        # V27(3): 지문이 같으면 해제는 필요 없다(동작은 같음)
        else:
            why = f"문서 지문 불일치(평가셋 {str(got['sha256'])[:12]} 대 우리 {spec['sha256'][:12]})"
            if spec.get("release"):                                                # V27(2): 지문 불일치는 지우지 않고 해제 기록을 함께 남긴다
                doc_status[did] = {"status": "released", "state": got["state"], "reason": why, "hash_mismatch": True, "released": dict(spec["release"])}
            else:
                doc_status[did] = {"status": "hold", "state": got["state"], "reason": why, "hash_mismatch": True}

    # 질문 정답표
    questions, by_id, row_kinds, states = [], {}, {}, {"review": {}, "kind": {}}
    for r in wb[sh["questions"]].iter_rows(min_row=first):
        cells = [c.value for c in r]
        kind = cells[qi["row_kind"]]
        label = "(빈칸)" if _blank(kind) else str(kind).strip()
        row_kinds[label] = row_kinds.get(label, 0) + 1
        read = [cells[qi[k]] for k in ("id", "main_type", "sub_type", "text", "answerable")]
        where = f"{sh['questions']} {r[0].row}행"
        if label != prof["row_kind_read"]:
            if any(not _blank(v) for v in read):                     # V2: 읽는 열(자동 열 제외)에 값이 있으면 조용히 넘기지 않는다
                raise XlsxGoldError(f"{where}: 행 구분이 {label!r}(읽는 값 {prof['row_kind_read']!r}이 아님)인데 질문 칸에 값이 있습니다")
            row_kinds[label] = row_kinds.get(label, 0)
            continue
        qid = cells[qi["id"]]
        if _blank(qid) or not isinstance(qid, str):
            raise XlsxGoldError(f"{where}: 질문 ID가 비었거나 글이 아닙니다")
        qid, where = qid.strip(), f"{sh['questions']} {r[0].row}행({qid.strip()})"
        if qid in by_id:
            raise XlsxGoldError(f"{where}: 질문 ID 중복")
        main = _text(cells[qi["main_type"]], where, "주유형", required=True)
        sub = _text(cells[qi["sub_type"]], where, "신규 하위 유형", required=True)
        qtype = main if sub == prof["sub_type_none"] else f"{main}{fmt['group_separator']}{sub}"
        if qtype not in prof["types"]:
            raise XlsxGoldError(f"{where}: 유형 {qtype!r}이 설정의 허용 목록 {prof['types']}에 없습니다")
        if qtype not in fmt["types"]:
            raise XlsxGoldError(f"{where}: 유형 {qtype!r}이 gold_format.yaml의 types에 없습니다(채점기가 읽지 못함)")
        text = _text(cells[qi["text"]], where, "사용자 질문", required=True)
        av = _text(cells[qi["answerable"]], where, "답변 가능 범위", required=True)
        if av not in prof["answerable_values"]:
            raise XlsxGoldError(f"{where}: 답변 가능 범위 값 {av!r}을 모릅니다(설정에 있는 값 {list(prof['answerable_values'])})")
        autos = {}
        for k, name in (("auto_required", "기본 필수 단위"), ("auto_conditional", "조건부 필수 단위"), ("auto_optional", "보조 단위")):
            v = cells[qi[k]]
            if v is None:                                            # V21: 계산된 값이 저장되지 않은 파일이면 조용히 건너뛰지 않는다
                raise XlsxGoldError(f"{where}: {name}(자동 열)의 저장된 값이 없습니다(엑셀에서 열어 저장한 파일을 쓰세요)")
            autos[k] = _whole(v, where, name)
        for key, bucket in (("state_review", "review"), ("state_kind", "kind")):
            v = cells[qi[key]]
            sv = "(빈칸)" if _blank(v) else str(v).strip()
            states[bucket][sv] = states[bucket].get(sv, 0) + 1
        d = {"case_id": qid, "type": qtype, "request": text, "answerable": bool(prof["answerable_values"][av]), "gold_evidence": [],
             "slot_rule": scoring.UNIT_BUNDLE}
        by_id[qid] = {"d": d, "autos": autos, "units": {}}
        questions.append(d)
    if not questions:
        raise XlsxGoldError(f"{sh['questions']}에 읽을 질문(행 구분 {prof['row_kind_read']!r})이 하나도 없습니다")

    # 질문별 근거
    n_evidence, blank_rows = 0, 0
    for r in wb[sh["evidence"]].iter_rows(min_row=first):
        cells = [c.value for c in r]
        qid = cells[ei["question_id"]]
        read = [cells[ei[k]] for k in ("unit", "bundle", "role", "document", "sheet", "cells", "area", "condition")]
        if _blank(qid):
            if any(not _blank(v) for v in read):                      # V2: 자동 열은 보지 않는다
                raise XlsxGoldError(f"{sh['evidence']} {r[0].row}행: 질문 ID가 비었는데 읽는 열에 값이 있습니다")
            blank_rows += 1
            continue
        qid = str(qid).strip()
        where = f"{sh['evidence']} {r[0].row}행({qid})"
        if qid not in by_id:
            raise XlsxGoldError(f"{where}: 질문 정답표에 없는 질문 ID(행 구분이 읽는 값이 아니거나 오타)")
        unit = _text(cells[ei["unit"]], where, "채점 단위 ID", required=True)
        bundle = _text(cells[ei["bundle"]], where, "대체 묶음 ID", required=True)
        role = _text(cells[ei["role"]], where, "근거 등급", required=True)
        if role not in prof["role_values"]:                           # V9
            raise XlsxGoldError(f"{where}: 근거 등급 값 {role!r}을 모릅니다(설정에 있는 값 {list(prof['role_values'])})")
        cond = _text(cells[ei["condition"]], where, "적용 조건", required=True)
        if cond != prof["condition_expected"]:                        # V3: 읽지 않지만 조건부 근거가 생기면 조용히 무시하지 않는다
            raise XlsxGoldError(f"{where}: 적용 조건이 {prof['condition_expected']!r}이 아닙니다: {cond!r}(이 설정은 조건부 채점을 읽지 않습니다)")
        did = _text(cells[ei["document"]], where, "문서 ID", required=True)
        if did not in prof["documents"]:                              # V4
            raise XlsxGoldError(f"{where}: 문서 ID {did!r}가 설정의 documents에 없습니다")
        if doc_status[did]["status"] == "unverifiable":
            raise XlsxGoldError(f"{where}: 문서 {did}의 지문을 평가셋 '{sh['documents']}' 시트에서 확인할 수 없습니다")
        sheet = cells[ei["sheet"]]                                    # 이슈 26: 시트 이름은 글자 하나도 바꾸지 않는다(끝 공백이 있는 실제 시트가 있다)
        if not isinstance(sheet, str) or not sheet.strip():
            raise XlsxGoldError(f"{where}: 시트 칸이 비었거나 글이 아닙니다: {sheet!r}")
        toks = _cell_list(cells[ei["cells"]], where)
        area = _text(cells[ei["area"]], where, "검토 영역")
        ev = {"file": prof["documents"][did]["file"], "sheet": sheet, "cells": toks, "required": bool(prof["role_values"][role]),
              "unit": unit, "bundle": bundle}
        if area is not None:
            ev["area"] = area
        q = by_id[qid]
        q["d"]["gold_evidence"].append(ev)
        q["units"].setdefault(unit, []).append((bundle, ev["required"], did))
        n_evidence += 1

    # V12, V15: 단위 안 역할 혼재 검사와 자동 열 대조(보류를 적용하기 전에 전체로 먼저)
    req_total = opt_total = 0
    shapes = {"단일 행": 0, "여러 행·묶음 1종": 0, "여러 행·묶음 여러 종": 0}
    for qid, q in by_id.items():
        req_units = opt_units = 0
        for unit, rows in q["units"].items():
            if len({r[1] for r in rows}) > 1:
                raise XlsxGoldError(f"질문 {qid}: 채점 단위 {unit!r} 안에 필수와 보조가 섞여 있습니다")
            if rows[0][1]:
                req_units += 1
            else:
                opt_units += 1
            nb = len({r[0] for r in rows})
            shapes["단일 행" if len(rows) == 1 else "여러 행·묶음 1종" if nb == 1 else "여러 행·묶음 여러 종"] += 1
        a = q["autos"]
        if a["auto_conditional"] != 0:
            raise XlsxGoldError(f"질문 {qid}: 조건부 필수 단위가 {a['auto_conditional']}개입니다(이 설정은 조건부 채점을 읽지 않습니다)")
        if (req_units, opt_units) != (a["auto_required"], a["auto_optional"]):
            raise XlsxGoldError(f"질문 {qid}: 근거 시트에서 센 필수·보조 단위 {req_units}·{opt_units}가 정답표의 자동 열 "
                                f"{a['auto_required']}·{a['auto_optional']}과 다릅니다")
        req_total += req_units
        opt_total += opt_units

    # V5: 보류는 질문 단위로 뺀다
    held, kept = [], []
    for d in questions:
        docs = {}
        for rows in by_id[d["case_id"]]["units"].values():
            for _b, role, did in rows:
                if doc_status[did]["status"] == "hold":
                    docs.setdefault(did, set()).add("필수" if role else "보조")
        if docs:
            held.append({"id": d["case_id"], "documents": sorted(docs), "evidence_roles": {k: sorted(v) for k, v in sorted(docs.items())},
                         "reason": "; ".join(f"{k}: {doc_status[k]['reason']}" for k in sorted(docs))})
        else:
            kept.append(d)
    if not kept:
        raise XlsxGoldError("보류를 빼면 남는 질문이 없습니다")
    try:
        scoring.parse_gold(kept)                                      # 채점기의 규칙(필수 칸 1~8개, 단위·묶음 검사 등)을 그대로 통과해야 한다
    except scoring.ScoringError as err:
        raise XlsxGoldError(f"채점기 검사에서 멈춤: {err}") from None
    by_type = lambda qs: dict(sorted({t: sum(1 for d in qs if d["type"] == t) for t in {d["type"] for d in qs}}.items()))
    profile_report = {"name": prof["profile"], "rule": prof["rule"], "documents": doc_status, "held_questions": held,
                      "n_questions_before_hold": len(questions), "n_evidence_before_hold": n_evidence,
                      "units_before_hold": {"required": req_total, "optional": opt_total}, "unit_shapes": shapes,
                      "row_kinds": row_kinds, "blank_evidence_rows": blank_rows,
                      "by_type_before_hold": by_type(questions), "by_type_after_hold": by_type(kept),
                      "source_state": {"정답 검수 상태": states["review"], "평가 구분": states["kind"],
                                       "문서 현행 확인 상태": {k: v["state"] for k, v in doc_status.items()}}}
    return kept, {"unused_columns": {sh["questions"]: sorted(q_unused), sh["evidence"]: sorted(e_unused), sh["documents"]: sorted(d_unused)},
                  "profile": profile_report}


# ---------- 묶음 ㉖: 공개 문서 평가셋 v0.2 읽기(기준은 CHANGES.md의 W1~W12) ----------
# v0.5 읽기 함수(read_gold_profile)는 건드리지 않고 이 함수를 따로 둔다. 공통 보조 함수만 같이 쓴다.
_V02_KEYS = {"profile", "rule", "sheets", "header_row", "first_row", "question_columns", "question_list_columns", "evidence_columns",
             "document_columns", "role_values", "type_map", "documents", "expected_counts"}
_V02_QCOLS = ("id", "type", "text", "state")
_V02_LCOLS = ("id", "text")
_V02_ECOLS = ("question_id", "unit", "bundle", "role", "file", "sheet_cells")
_V02_DCOLS = ("id", "file", "sha256", "state")
_V02_COUNTS = ("questions", "evidence", "required_rows", "required_units")
_PROFILE_NAMES = ("v0.5", "v0.2", "b-v0.1", "c-v0.1", "e-v0.2", "d-v0.1")


def profile_name(path: str):
    """읽기 설정 파일의 profile 이름. 읽을 수 없으면 None(호출한 쪽이 멈춘다)."""
    import yaml

    with open(path, encoding="utf-8") as f:
        raw = yaml.safe_load(f)
    return raw.get("profile") if isinstance(raw, dict) else None


def load_profile_v02(path: str) -> dict:
    import yaml

    with open(path, encoding="utf-8") as f:
        raw = yaml.safe_load(f)
    if not isinstance(raw, dict):
        raise XlsxGoldError(f"{path}: 설정은 키와 값의 목록이어야 합니다")
    extra, missing = set(raw) - _V02_KEYS, _V02_KEYS - set(raw)
    if extra or missing:
        raise XlsxGoldError(f"{path}: 칸이 맞지 않습니다(모르는 칸 {sorted(extra)}, 빠진 칸 {sorted(missing)})")
    if raw["profile"] != "v0.2":
        raise XlsxGoldError(f"{path}: profile은 'v0.2'여야 합니다: {raw['profile']!r}")
    if raw["rule"] != scoring.UNIT_BUNDLE:
        raise XlsxGoldError(f"{path}: rule은 {scoring.UNIT_BUNDLE!r}만 쓸 수 있습니다: {raw['rule']!r}")
    for key, names in (("question_columns", _V02_QCOLS), ("question_list_columns", _V02_LCOLS), ("evidence_columns", _V02_ECOLS),
                       ("document_columns", _V02_DCOLS)):
        got = raw[key]
        if not isinstance(got, dict) or set(got) != set(names) or not all(isinstance(v, str) and v.strip() for v in got.values()):
            raise XlsxGoldError(f"{path}: {key}는 {list(names)}를 모두 가진 비어 있지 않은 문자열이어야 합니다")
    if not isinstance(raw["sheets"], dict) or set(raw["sheets"]) != {"questions", "question_list", "evidence", "documents"} \
            or not all(isinstance(v, str) and v for v in raw["sheets"].values()):
        raise XlsxGoldError(f"{path}: sheets는 questions, question_list, evidence, documents 이름을 가져야 합니다")
    for k in ("header_row", "first_row"):
        if isinstance(raw[k], bool) or not isinstance(raw[k], int) or raw[k] < 1:
            raise XlsxGoldError(f"{path}: {k}는 1 이상의 정수여야 합니다")
    if not isinstance(raw["role_values"], dict) or not raw["role_values"] or not all(isinstance(b, bool) for b in raw["role_values"].values()):
        raise XlsxGoldError(f"{path}: role_values는 값 이름 -> true/false 목록이어야 합니다")
    tm = raw["type_map"]
    if not isinstance(tm, dict) or not tm or not all(isinstance(k, str) and isinstance(v, str) and k.strip() and v.strip() for k, v in tm.items()):
        raise XlsxGoldError(f"{path}: type_map은 원본 유형 이름 -> 채점기 유형 이름 목록이어야 합니다")
    ec = raw["expected_counts"]
    if not isinstance(ec, dict) or set(ec) != set(_V02_COUNTS) or not all(isinstance(v, int) and not isinstance(v, bool) and v >= 0 for v in ec.values()):
        raise XlsxGoldError(f"{path}: expected_counts는 {list(_V02_COUNTS)}를 모두 가진 0 이상의 정수여야 합니다")
    if not isinstance(raw["documents"], dict) or not raw["documents"]:
        raise XlsxGoldError(f"{path}: documents가 비어 있습니다")
    for doc, v in raw["documents"].items():
        if not isinstance(v, dict) or not {"file", "sha256"} <= set(v) or set(v) - {"file", "sha256", "release"} \
                or not all(isinstance(v[k], str) and v[k].strip() for k in ("file", "sha256")):
            raise XlsxGoldError(f"{path}: documents.{doc}는 file과 sha256(선택으로 release)을 가져야 합니다")
        _file_name(v["file"], f"{path} documents.{doc}")
        rel = v.get("release")
        if "release" in v and (not isinstance(rel, dict) or set(rel) != {"by", "date", "reason"}
                               or not all(isinstance(x, str) and x.strip() for x in rel.values())):
            raise XlsxGoldError(f"{path}: documents.{doc}.release는 by, date, reason을 모두 가진 비어 있지 않은 글이어야 합니다")
    return raw


def read_gold_profile_v02(path: str, profile_path: str):
    """공개 문서 평가셋 v0.2(전문가 네트워크)를 설정대로 읽는다. (정답표 목록, 변환 보고)를 돌려주고 어긋나면 XlsxGoldError.
    열이 없는 값(답변 가능 여부, 적용 조건, 검토 영역)은 만들어 채우지 않고 보고에 '열 없음'으로 남긴다(W1·W3·W6)."""
    import openpyxl

    prof = load_profile_v02(profile_path)
    fmt = scoring.load_format()
    wb = openpyxl.load_workbook(path, data_only=True)
    sh = prof["sheets"]
    for s_ in sh.values():
        if s_ not in wb.sheetnames:
            raise XlsxGoldError(f"시트가 없습니다: {s_}")
    hdr, first = prof["header_row"], prof["first_row"]
    qi, q_unused = _profile_header(wb[sh["questions"]], hdr, prof["question_columns"], sh["questions"])
    li, l_unused = _profile_header(wb[sh["question_list"]], hdr, prof["question_list_columns"], sh["question_list"])
    ei, e_unused = _profile_header(wb[sh["evidence"]], hdr, prof["evidence_columns"], sh["evidence"])
    di, d_unused = _profile_header(wb[sh["documents"]], hdr, prof["document_columns"], sh["documents"])

    # W7: 기준 문서 시트(파일명 -> 문서 ID)와 지문. 상태는 코드가 계산한다
    sheet_docs, file_to_doc = {}, {}
    for r in wb[sh["documents"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        did = c[di["id"]]
        if _blank(did):
            continue
        did = str(did).strip()
        fname = c[di["file"]]
        if _blank(fname) or did in sheet_docs:
            raise XlsxGoldError(f"{sh['documents']} {r[0].row}행: 파일명이 비었거나 문서 ID({did})가 중복입니다")
        fname = _file_name(str(fname), f"{sh['documents']} {r[0].row}행")
        if fname in file_to_doc:
            raise XlsxGoldError(f"{sh['documents']} {r[0].row}행: 파일명 {fname!r}이 둘 이상의 문서 ID에 있습니다")
        file_to_doc[fname] = did
        sheet_docs[did] = {"file": fname, "sha256": None if _blank(c[di["sha256"]]) else str(c[di["sha256"]]).strip().lower(),
                           "state": None if _blank(c[di["state"]]) else str(c[di["state"]]).strip()}
    doc_status = {}
    for did, spec in prof["documents"].items():
        got = sheet_docs.get(did)
        if got is None:
            doc_status[did] = {"status": "unverifiable", "state": None, "reason": "평가셋 기준 문서 시트에 없음"}
        elif got["file"] != spec["file"]:
            raise XlsxGoldError(f"문서 {did}: 설정의 파일 이름 {spec['file']!r}이 평가셋 기준 문서 시트의 {got['file']!r}과 다릅니다")
        elif got["sha256"] == spec["sha256"].lower():
            doc_status[did] = {"status": "ok", "state": got["state"], "reason": None, "hash_mismatch": False}
            if spec.get("release"):
                doc_status[did]["released_unneeded"] = dict(spec["release"])
        else:
            why = f"문서 지문 불일치(평가셋 {str(got['sha256'])[:12]} 대 우리 {spec['sha256'][:12]})"
            if spec.get("release"):
                doc_status[did] = {"status": "released", "state": got["state"], "reason": why, "hash_mismatch": True, "released": dict(spec["release"])}
            else:
                doc_status[did] = {"status": "hold", "state": got["state"], "reason": why, "hash_mismatch": True}

    # W3·W4: 질문 정답표
    questions, by_id, states = [], {}, {}
    for r in wb[sh["questions"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        qid = c[qi["id"]]
        where = f"{sh['questions']} {r[0].row}행"
        if _blank(qid):
            if any(not _blank(c[qi[k]]) for k in ("type", "text")):
                raise XlsxGoldError(f"{where}: 질문 ID가 비었는데 읽는 열에 값이 있습니다")
            continue
        if not isinstance(qid, str):
            raise XlsxGoldError(f"{where}: 질문 ID가 글이 아닙니다")
        qid = qid.strip()
        where = f"{where}({qid})"
        if qid in by_id:
            raise XlsxGoldError(f"{where}: 질문 ID 중복")
        src_type = _text(c[qi["type"]], where, "유형", required=True)
        if src_type not in prof["type_map"]:
            raise XlsxGoldError(f"{where}: 유형 {src_type!r}이 설정의 type_map에 없습니다(있는 이름 {sorted(prof['type_map'])})")
        qtype = prof["type_map"][src_type]
        if qtype not in fmt["types"]:
            raise XlsxGoldError(f"{where}: 유형 {qtype!r}이 gold_format.yaml의 types에 없습니다(채점기가 읽지 못함)")
        text = _text(c[qi["text"]], where, "질문", required=True)
        sv = c[qi["state"]]
        sv = "(빈칸)" if _blank(sv) else str(sv).strip()
        states[sv] = states.get(sv, 0) + 1
        d = {"case_id": qid, "type": qtype, "request": text, "gold_evidence": [], "slot_rule": scoring.UNIT_BUNDLE}   # answerable은 정하지 않는다(W3)
        by_id[qid] = {"d": d, "units": {}}
        questions.append(d)
    if not questions:
        raise XlsxGoldError(f"{sh['questions']}에 읽을 질문이 하나도 없습니다")

    # W3: 질문 목록 시트와 같은지 대조
    listed = {}
    for r in wb[sh["question_list"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        if _blank(c[li["id"]]):
            continue
        qid = str(c[li["id"]]).strip()
        if qid in listed:
            raise XlsxGoldError(f"{sh['question_list']} {r[0].row}행: 질문 ID {qid} 중복")
        listed[qid] = None if _blank(c[li["text"]]) else str(c[li["text"]]).strip()
    if set(listed) != set(by_id):
        raise XlsxGoldError(f"질문 목록과 질문 정답표의 질문 ID가 다릅니다: 목록에만 {sorted(set(listed) - set(by_id))}, "
                            f"정답표에만 {sorted(set(by_id) - set(listed))}")
    diff_text = sorted(q for q in by_id if listed[q] != by_id[q]["d"]["request"])
    if diff_text:
        raise XlsxGoldError(f"질문 목록과 질문 정답표의 질문 글이 다른 질문이 있습니다: {diff_text}")

    # W5·W6: 질문별 근거
    n_evidence, blank_rows = 0, 0
    for r in wb[sh["evidence"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        qid = c[ei["question_id"]]
        read = [c[ei[k]] for k in ("unit", "bundle", "role", "file", "sheet_cells")]
        if _blank(qid):
            if any(not _blank(v) for v in read):
                raise XlsxGoldError(f"{sh['evidence']} {r[0].row}행: 질문 ID가 비었는데 읽는 열에 값이 있습니다")
            blank_rows += 1
            continue
        qid = str(qid).strip()
        where = f"{sh['evidence']} {r[0].row}행({qid})"
        if qid not in by_id:
            raise XlsxGoldError(f"{where}: 질문 정답표에 없는 질문 ID")
        unit = _text(c[ei["unit"]], where, "채점 단위 ID", required=True)
        bundle = _text(c[ei["bundle"]], where, "대체 묶음", required=True)
        role = _text(c[ei["role"]], where, "필수/보조", required=True)
        if role not in prof["role_values"]:
            raise XlsxGoldError(f"{where}: 필수/보조 값 {role!r}을 모릅니다(설정에 있는 값 {list(prof['role_values'])})")
        fname = _file_name(_text(c[ei["file"]], where, "파일명", required=True), where)
        if fname not in file_to_doc:
            raise XlsxGoldError(f"{where}: 파일명 {fname!r}이 평가셋 '{sh['documents']}' 시트에 없습니다")
        did = file_to_doc[fname]
        if did not in prof["documents"]:
            raise XlsxGoldError(f"{where}: 문서 ID {did!r}가 설정의 documents에 없습니다")
        if doc_status[did]["status"] == "unverifiable":
            raise XlsxGoldError(f"{where}: 문서 {did}의 지문을 평가셋 '{sh['documents']}' 시트에서 확인할 수 없습니다")
        sc = c[ei["sheet_cells"]]
        if not isinstance(sc, str) or "!" not in sc:
            raise XlsxGoldError(f"{where}: 시트·셀 칸이 '시트!셀범위' 모양이 아닙니다: {sc!r}")
        sheet, cells = sc.rsplit("!", 1)                              # W5: 마지막 !에서 나눈다. 시트 이름은 글자 하나도 바꾸지 않는다(이슈 26)
        if not sheet.strip():
            raise XlsxGoldError(f"{where}: 시트 이름이 비었습니다: {sc!r}")
        toks = _cell_list(cells, where)
        ev = {"file": fname, "sheet": sheet, "cells": toks, "required": bool(prof["role_values"][role]), "unit": unit, "bundle": bundle}
        q = by_id[qid]
        q["d"]["gold_evidence"].append(ev)
        q["units"].setdefault(unit, []).append((bundle, ev["required"], did))
        n_evidence += 1

    req_rows = req_units = opt_units = 0
    shapes = {"단일 행": 0, "여러 행·묶음 1종": 0, "여러 행·묶음 여러 종": 0}
    for qid, q in by_id.items():
        for unit, rows in q["units"].items():
            if len({x[1] for x in rows}) > 1:
                raise XlsxGoldError(f"질문 {qid}: 채점 단위 {unit!r} 안에 필수와 보조가 섞여 있습니다")
            if rows[0][1]:
                req_units += 1
                req_rows += len(rows)
            else:
                opt_units += 1
            nb = len({x[0] for x in rows})
            shapes["단일 행" if len(rows) == 1 else "여러 행·묶음 1종" if nb == 1 else "여러 행·묶음 여러 종"] += 1
    # W8: 설정에 적은 기대 개수와 대조
    got_counts = {"questions": len(questions), "evidence": n_evidence, "required_rows": req_rows, "required_units": req_units}
    bad = {k: (got_counts[k], prof["expected_counts"][k]) for k in _V02_COUNTS if got_counts[k] != prof["expected_counts"][k]}
    if bad:
        raise XlsxGoldError("읽은 개수가 설정의 expected_counts와 다릅니다(읽은 값, 기대 값): " + json.dumps(bad, ensure_ascii=False))

    held, kept = [], []
    for d in questions:
        docs = {}
        for rows in by_id[d["case_id"]]["units"].values():
            for _b, role_, did in rows:
                if doc_status[did]["status"] == "hold":
                    docs.setdefault(did, set()).add("필수" if role_ else "보조")
        if docs:
            held.append({"id": d["case_id"], "documents": sorted(docs), "evidence_roles": {k: sorted(v) for k, v in sorted(docs.items())},
                         "reason": "; ".join(f"{k}: {doc_status[k]['reason']}" for k in sorted(docs))})
        else:
            kept.append(d)
    if not kept:
        raise XlsxGoldError("보류를 빼면 남는 질문이 없습니다")
    try:
        scoring.parse_gold(kept)
    except scoring.ScoringError as err:
        raise XlsxGoldError(f"채점기 검사에서 멈춤: {err}") from None
    by_type = lambda qs: dict(sorted({t: sum(1 for d in qs if d["type"] == t) for t in {d["type"] for d in qs}}.items()))
    profile_report = {"name": prof["profile"], "rule": prof["rule"], "documents": doc_status, "held_questions": held,
                      "n_questions_before_hold": len(questions), "n_evidence_before_hold": n_evidence,
                      "units_before_hold": {"required": req_units, "optional": opt_units}, "unit_shapes": shapes,
                      "row_kinds": {}, "blank_evidence_rows": blank_rows, "required_rows": req_rows,
                      "by_type_before_hold": by_type(questions), "by_type_after_hold": by_type(kept),
                      "not_available": {"답변 가능 범위": "열 없음, 채점기 기본값(가능)을 씀", "적용 조건": "열 없음, 조건부 채점을 하지 않음",
                                        "검토 영역": "열 없음, 영역별 누락 보고 없음"},
                      "source_state": {"상태": states, "문서 출처 상태": {k: v["state"] for k, v in doc_status.items()}}}
    return kept, {"unused_columns": {sh["questions"]: sorted(q_unused), sh["question_list"]: sorted(l_unused), sh["evidence"]: sorted(e_unused),
                                     sh["documents"]: sorted(d_unused)},
                  "profile": profile_report}


# ---------- 묶음 72: 서비스 B 평가셋 v0.1 읽기(기준은 CHANGES.md 묶음 72 D1) ----------
# v0.2 읽기 함수는 건드리지 않고 이 함수를 따로 둔다. 다른 점: 대체 묶음이 전부 빈칸(단위마다 묶음 하나), 유형 이름 그대로,
# 쟁점 질문을 빼거나(main) 그것만 읽는다(issues). 기대 개수는 빼기 전 개수로 대조한다.
_FMTB_KEYS = _V02_KEYS | {"bundle_mode", "issue_questions", "question_subset"}
_FMTB_BUNDLE_MODES = ("blank_single",)
_FMTB_SUBSETS = ("main", "issues", "all")                       # all: 묶음 76 A안(25문항 전부)
FMTB_SINGLE_BUNDLE = "(묶음 없음)"                                   # 빈 대체 묶음을 읽은 값. 단위 안 근거는 모두 필요(AND)


def load_profile_fmtb01(path: str) -> dict:
    import yaml

    with open(path, encoding="utf-8") as f:
        raw = yaml.safe_load(f)
    if not isinstance(raw, dict):
        raise XlsxGoldError(f"{path}: 설정은 키와 값의 목록이어야 합니다")
    extra, missing = set(raw) - _FMTB_KEYS, _FMTB_KEYS - set(raw)
    if extra or missing:
        raise XlsxGoldError(f"{path}: 칸이 맞지 않습니다(모르는 칸 {sorted(extra)}, 빠진 칸 {sorted(missing)})")
    if raw["profile"] != "b-v0.1":
        raise XlsxGoldError(f"{path}: profile은 'b-v0.1'이어야 합니다: {raw['profile']!r}")
    if raw["rule"] != scoring.UNIT_BUNDLE:
        raise XlsxGoldError(f"{path}: rule은 {scoring.UNIT_BUNDLE!r}만 쓸 수 있습니다: {raw['rule']!r}")
    for key, names in (("question_columns", _V02_QCOLS), ("question_list_columns", _V02_LCOLS), ("evidence_columns", _V02_ECOLS),
                       ("document_columns", _V02_DCOLS)):
        got = raw[key]
        if not isinstance(got, dict) or set(got) != set(names) or not all(isinstance(v, str) and v.strip() for v in got.values()):
            raise XlsxGoldError(f"{path}: {key}는 {list(names)}를 모두 가진 비어 있지 않은 문자열이어야 합니다")
    if not isinstance(raw["sheets"], dict) or set(raw["sheets"]) != {"questions", "question_list", "evidence", "documents"} \
            or not all(isinstance(v, str) and v for v in raw["sheets"].values()):
        raise XlsxGoldError(f"{path}: sheets는 questions, question_list, evidence, documents 이름을 가져야 합니다")
    for k in ("header_row", "first_row"):
        if isinstance(raw[k], bool) or not isinstance(raw[k], int) or raw[k] < 1:
            raise XlsxGoldError(f"{path}: {k}는 1 이상의 정수여야 합니다")
    if not isinstance(raw["role_values"], dict) or not raw["role_values"] or not all(isinstance(b, bool) for b in raw["role_values"].values()):
        raise XlsxGoldError(f"{path}: role_values는 값 이름 -> true/false 목록이어야 합니다")
    tm = raw["type_map"]
    if not isinstance(tm, dict) or not tm or not all(isinstance(k, str) and isinstance(v, str) and k.strip() and v.strip() for k, v in tm.items()):
        raise XlsxGoldError(f"{path}: type_map은 원본 유형 이름 -> 채점기 유형 이름 목록이어야 합니다")
    ec = raw["expected_counts"]
    if not isinstance(ec, dict) or set(ec) != set(_V02_COUNTS) or not all(isinstance(v, int) and not isinstance(v, bool) and v >= 0 for v in ec.values()):
        raise XlsxGoldError(f"{path}: expected_counts는 {list(_V02_COUNTS)}를 모두 가진 0 이상의 정수여야 합니다")
    if not isinstance(raw["documents"], dict) or not raw["documents"]:
        raise XlsxGoldError(f"{path}: documents가 비어 있습니다")
    for doc, v in raw["documents"].items():
        if not isinstance(v, dict) or set(v) != {"file", "sha256"} or not all(isinstance(v[k], str) and v[k].strip() for k in ("file", "sha256")):
            raise XlsxGoldError(f"{path}: documents.{doc}는 file과 sha256만 가져야 합니다")
        _file_name(v["file"], f"{path} documents.{doc}")
    if raw["bundle_mode"] not in _FMTB_BUNDLE_MODES:
        raise XlsxGoldError(f"{path}: bundle_mode는 {list(_FMTB_BUNDLE_MODES)} 중 하나여야 합니다: {raw['bundle_mode']!r}")
    if raw["question_subset"] not in _FMTB_SUBSETS:
        raise XlsxGoldError(f"{path}: question_subset은 {list(_FMTB_SUBSETS)} 중 하나여야 합니다: {raw['question_subset']!r}")
    iq = raw["issue_questions"]
    if not isinstance(iq, dict) or set(iq) != {"ids", "reason"}:
        raise XlsxGoldError(f"{path}: issue_questions는 ids와 reason만 가져야 합니다")
    ids = iq["ids"]
    if not isinstance(ids, list) or not ids or not all(isinstance(x, str) and x.strip() for x in ids):
        raise XlsxGoldError(f"{path}: issue_questions.ids는 비어 있지 않은 질문 ID 목록이어야 합니다")
    dup = sorted({x for x in ids if ids.count(x) > 1})
    if dup:
        raise XlsxGoldError(f"{path}: issue_questions.ids에 같은 ID가 두 번 이상 있습니다: {dup}")
    if not isinstance(iq["reason"], str) or not iq["reason"].strip():
        raise XlsxGoldError(f"{path}: issue_questions.reason이 비었습니다")
    return raw


def read_gold_profile_fmtb01(path: str, profile_path: str):
    """서비스 B 평가셋 v0.1을 설정대로 읽는다. (정답표 목록, 변환 보고)를 돌려주고 어긋나면 XlsxGoldError(묶음 72 D1)."""
    import openpyxl

    prof = load_profile_fmtb01(profile_path)
    fmt = scoring.load_format()
    wb = openpyxl.load_workbook(path, data_only=True)
    sh = prof["sheets"]
    for s_ in sh.values():
        if s_ not in wb.sheetnames:
            raise XlsxGoldError(f"시트가 없습니다: {s_}")
    hdr, first = prof["header_row"], prof["first_row"]
    qi, q_unused = _profile_header(wb[sh["questions"]], hdr, prof["question_columns"], sh["questions"])
    li, l_unused = _profile_header(wb[sh["question_list"]], hdr, prof["question_list_columns"], sh["question_list"])
    ei, e_unused = _profile_header(wb[sh["evidence"]], hdr, prof["evidence_columns"], sh["evidence"])
    di, d_unused = _profile_header(wb[sh["documents"]], hdr, prof["document_columns"], sh["documents"])

    # 기준 문서 시트(파일명 -> 문서 ID)와 지문. 상태는 코드가 계산한다(v0.2와 같음, 풀어 주기 칸은 없음)
    sheet_docs, file_to_doc = {}, {}
    for r in wb[sh["documents"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        did = c[di["id"]]
        if _blank(did):
            continue
        did = str(did).strip()
        fname = c[di["file"]]
        if _blank(fname) or did in sheet_docs:
            raise XlsxGoldError(f"{sh['documents']} {r[0].row}행: 파일명이 비었거나 문서 ID({did})가 중복입니다")
        fname = _file_name(str(fname), f"{sh['documents']} {r[0].row}행")
        if fname in file_to_doc:
            raise XlsxGoldError(f"{sh['documents']} {r[0].row}행: 파일명 {fname!r}이 둘 이상의 문서 ID에 있습니다")
        file_to_doc[fname] = did
        sheet_docs[did] = {"file": fname, "sha256": None if _blank(c[di["sha256"]]) else str(c[di["sha256"]]).strip().lower(),
                           "state": None if _blank(c[di["state"]]) else str(c[di["state"]]).strip()}
    doc_status = {}
    for did, spec in prof["documents"].items():
        got = sheet_docs.get(did)
        if got is None:
            doc_status[did] = {"status": "unverifiable", "state": None, "reason": "평가셋 기준 문서 시트에 없음"}
        elif got["file"] != spec["file"]:
            raise XlsxGoldError(f"문서 {did}: 설정의 파일 이름 {spec['file']!r}이 평가셋 기준 문서 시트의 {got['file']!r}과 다릅니다")
        elif got["sha256"] == spec["sha256"].lower():
            doc_status[did] = {"status": "ok", "state": got["state"], "reason": None, "hash_mismatch": False}
        else:
            doc_status[did] = {"status": "hold", "state": got["state"], "hash_mismatch": True,
                               "reason": f"문서 지문 불일치(평가셋 {str(got['sha256'])[:12]} 대 우리 {spec['sha256'][:12]})"}

    # 질문 정답표(유형은 type_map으로 바꾸고 채점기 목록에 있어야 한다)
    questions, by_id, states = [], {}, {}
    for r in wb[sh["questions"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        qid = c[qi["id"]]
        where = f"{sh['questions']} {r[0].row}행"
        if _blank(qid):
            if any(not _blank(c[qi[k]]) for k in ("type", "text")):
                raise XlsxGoldError(f"{where}: 질문 ID가 비었는데 읽는 열에 값이 있습니다")
            continue
        if not isinstance(qid, str):
            raise XlsxGoldError(f"{where}: 질문 ID가 글이 아닙니다")
        qid = qid.strip()
        where = f"{where}({qid})"
        if qid in by_id:
            raise XlsxGoldError(f"{where}: 질문 ID 중복")
        src_type = _text(c[qi["type"]], where, "유형", required=True)
        if src_type not in prof["type_map"]:
            raise XlsxGoldError(f"{where}: 유형 {src_type!r}이 설정의 type_map에 없습니다(있는 이름 {sorted(prof['type_map'])})")
        qtype = prof["type_map"][src_type]
        if qtype not in fmt["types"]:
            raise XlsxGoldError(f"{where}: 유형 {qtype!r}이 gold_format.yaml의 types에 없습니다(채점기가 읽지 못함)")
        text = _text(c[qi["text"]], where, "질문", required=True)
        sv = c[qi["state"]]
        sv = "(빈칸)" if _blank(sv) else str(sv).strip()
        states[sv] = states.get(sv, 0) + 1
        d = {"case_id": qid, "type": qtype, "request": text, "gold_evidence": [], "slot_rule": scoring.UNIT_BUNDLE}
        by_id[qid] = {"d": d, "units": {}}
        questions.append(d)
    if not questions:
        raise XlsxGoldError(f"{sh['questions']}에 읽을 질문이 하나도 없습니다")

    # 질문 목록 시트와 같은지 대조
    listed = {}
    for r in wb[sh["question_list"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        if _blank(c[li["id"]]):
            continue
        qid = str(c[li["id"]]).strip()
        if qid in listed:
            raise XlsxGoldError(f"{sh['question_list']} {r[0].row}행: 질문 ID {qid} 중복")
        listed[qid] = None if _blank(c[li["text"]]) else str(c[li["text"]]).strip()
    if set(listed) != set(by_id):
        raise XlsxGoldError(f"질문 목록과 질문 정답표의 질문 ID가 다릅니다: 목록에만 {sorted(set(listed) - set(by_id))}, "
                            f"정답표에만 {sorted(set(by_id) - set(listed))}")
    diff_text = sorted(q for q in by_id if listed[q] != by_id[q]["d"]["request"])
    if diff_text:
        raise XlsxGoldError(f"질문 목록과 질문 정답표의 질문 글이 다른 질문이 있습니다: {diff_text}")

    # 질문별 근거(D1a: 대체 묶음은 모두 빈칸이어야 하고 단위마다 묶음 하나)
    n_evidence, blank_rows = 0, 0
    for r in wb[sh["evidence"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        qid = c[ei["question_id"]]
        read = [c[ei[k]] for k in ("unit", "bundle", "role", "file", "sheet_cells")]
        if _blank(qid):
            if any(not _blank(v) for v in read):
                raise XlsxGoldError(f"{sh['evidence']} {r[0].row}행: 질문 ID가 비었는데 읽는 열에 값이 있습니다")
            blank_rows += 1
            continue
        qid = str(qid).strip()
        where = f"{sh['evidence']} {r[0].row}행({qid})"
        if qid not in by_id:
            raise XlsxGoldError(f"{where}: 질문 정답표에 없는 질문 ID")
        unit = _text(c[ei["unit"]], where, "채점 단위 ID", required=True)
        if not _blank(c[ei["bundle"]]):
            raise XlsxGoldError(f"{where}: bundle_mode {prof['bundle_mode']!r}인데 대체 묶음 칸에 값이 있습니다: {c[ei['bundle']]!r}")
        bundle = FMTB_SINGLE_BUNDLE
        role = _text(c[ei["role"]], where, "필수/보조", required=True)
        if role not in prof["role_values"]:
            raise XlsxGoldError(f"{where}: 필수/보조 값 {role!r}을 모릅니다(설정에 있는 값 {list(prof['role_values'])})")
        fname = _file_name(_text(c[ei["file"]], where, "파일명", required=True), where)
        if fname not in file_to_doc:
            raise XlsxGoldError(f"{where}: 파일명 {fname!r}이 평가셋 '{sh['documents']}' 시트에 없습니다")
        did = file_to_doc[fname]
        if did not in prof["documents"]:
            raise XlsxGoldError(f"{where}: 문서 ID {did!r}가 설정의 documents에 없습니다")
        if doc_status[did]["status"] == "unverifiable":
            raise XlsxGoldError(f"{where}: 문서 {did}의 지문을 평가셋 '{sh['documents']}' 시트에서 확인할 수 없습니다")
        sc = c[ei["sheet_cells"]]
        if not isinstance(sc, str) or "!" not in sc:
            raise XlsxGoldError(f"{where}: 시트·셀 칸이 '시트!셀범위' 모양이 아닙니다: {sc!r}")
        sheet, cells = sc.rsplit("!", 1)                              # 마지막 !에서 나눈다. 시트 이름은 글자 하나도 바꾸지 않는다(이슈 26)
        if not sheet.strip():
            raise XlsxGoldError(f"{where}: 시트 이름이 비었습니다: {sc!r}")
        toks = _cell_list(cells, where)
        ev = {"file": fname, "sheet": sheet, "cells": toks, "required": bool(prof["role_values"][role]), "unit": unit, "bundle": bundle}
        q = by_id[qid]
        q["d"]["gold_evidence"].append(ev)
        q["units"].setdefault(unit, []).append((bundle, ev["required"], did))
        n_evidence += 1

    def unit_counts(ids):
        rr = ru = ou = 0
        for qid in ids:
            for unit, rows in by_id[qid]["units"].items():
                if rows[0][1]:
                    ru += 1
                    rr += len(rows)
                else:
                    ou += 1
        return rr, ru, ou

    for qid, q in by_id.items():
        for unit, rows in q["units"].items():
            if len({x[1] for x in rows}) > 1:
                raise XlsxGoldError(f"질문 {qid}: 채점 단위 {unit!r} 안에 필수와 보조가 섞여 있습니다")
    req_rows, req_units, opt_units = unit_counts(by_id)
    shapes = {"단일 행": 0, "여러 행": 0}
    for q in by_id.values():
        for rows in q["units"].values():
            shapes["단일 행" if len(rows) == 1 else "여러 행"] += 1
    # D1d: 기대 개수는 빼기 전 개수로 대조한다
    got_counts = {"questions": len(questions), "evidence": n_evidence, "required_rows": req_rows, "required_units": req_units}
    bad = {k: (got_counts[k], prof["expected_counts"][k]) for k in _V02_COUNTS if got_counts[k] != prof["expected_counts"][k]}
    if bad:
        raise XlsxGoldError("읽은 개수가 설정의 expected_counts와 다릅니다(읽은 값, 기대 값): " + json.dumps(bad, ensure_ascii=False))

    # D1c: 쟁점 질문을 빼거나(main) 그것만 남긴다(issues)
    issue_ids = [x.strip() for x in prof["issue_questions"]["ids"]]
    unknown = sorted(set(issue_ids) - set(by_id))
    if unknown:
        raise XlsxGoldError(f"issue_questions.ids에 정답지에 없는 질문 ID가 있습니다: {unknown}")
    subset = prof["question_subset"]
    chosen = list(questions) if subset == "all" else [d for d in questions if (d["case_id"] in issue_ids) == (subset == "issues")]
    if not chosen:
        raise XlsxGoldError(f"question_subset {subset!r}로 고르면 남는 질문이 없습니다")
    left_out = [d["case_id"] for d in questions if d not in chosen]

    held, kept = [], []
    for d in chosen:
        docs = {}
        for rows in by_id[d["case_id"]]["units"].values():
            for _b, role_, did in rows:
                if doc_status[did]["status"] == "hold":
                    docs.setdefault(did, set()).add("필수" if role_ else "보조")
        if docs:
            held.append({"id": d["case_id"], "documents": sorted(docs), "evidence_roles": {k: sorted(v) for k, v in sorted(docs.items())},
                         "reason": "; ".join(f"{k}: {doc_status[k]['reason']}" for k in sorted(docs))})
        else:
            kept.append(d)
    if not kept:
        raise XlsxGoldError("보류를 빼면 남는 질문이 없습니다")
    try:
        scoring.parse_gold(kept)
    except scoring.ScoringError as err:
        raise XlsxGoldError(f"채점기 검사에서 멈춤: {err}") from None
    sub_rr, sub_ru, sub_ou = unit_counts([d["case_id"] for d in chosen])
    by_type = lambda qs: dict(sorted({t: sum(1 for d in qs if d["type"] == t) for t in {d["type"] for d in qs}}.items()))
    profile_report = {"name": prof["profile"], "rule": prof["rule"], "documents": doc_status, "held_questions": held,
                      "n_questions_before_hold": len(chosen), "n_evidence_before_hold": sum(len(d["gold_evidence"]) for d in chosen),
                      "units_before_hold": {"required": sub_ru, "optional": sub_ou}, "unit_shapes": shapes,
                      "row_kinds": {}, "blank_evidence_rows": blank_rows, "required_rows": sub_rr,
                      "by_type_before_hold": by_type(chosen), "by_type_after_hold": by_type(kept),
                      "bundle_mode": {"mode": prof["bundle_mode"], "read_as": FMTB_SINGLE_BUNDLE,
                                      "meaning": "대체 묶음이 모두 빈칸이라 단위마다 묶음 하나(단위 안 근거는 모두 필요)"},
                      "question_subset": {"subset": subset, "issue_ids": issue_ids, "reason": prof["issue_questions"]["reason"].strip(),
                                          "left_out": left_out, "before_subset": got_counts},
                      "not_available": {"답변 가능 범위": "열 없음, 채점기 기본값(가능)을 씀", "적용 조건": "열 없음, 조건부 채점을 하지 않음",
                                        "검토 영역": "열 없음, 영역별 누락 보고 없음"},
                      "source_state": {"상태": states, "문서 출처 상태": {k: v["state"] for k, v in doc_status.items()}}}
    return kept, {"unused_columns": {sh["questions"]: sorted(q_unused), sh["question_list"]: sorted(l_unused), sh["evidence"]: sorted(e_unused),
                                     sh["documents"]: sorted(d_unused)},
                  "profile": profile_report}



# ---------- 묶음 96: 서비스 E(서비스 E) 평가셋 v0.2 읽기(기준은 CHANGES.md 묶음 96 F4) ----------
# 서비스 B 함수(fmtb01)를 바탕으로 새로 둔다(기존 함수는 고치지 않음, 10/11 교훈). 다른 점 두 가지:
#   1) 대체 묶음 칸이 모두 채워져 있다. 값을 그대로 묶음 이름으로 쓴다
#      (채점기 규칙: 같은 단위 안에서 같은 묶음은 모두 필요(AND), 다른 묶음끼리는 하나만(OR)). 빈칸이면 멈춘다.
#   2) 쟁점 문항이 없어 30문항을 모두 읽는다(issue_questions·question_subset 칸이 없다).
_NET_KEYS = _V02_KEYS | {"bundle_mode"}
_NET_BUNDLE_MODES = ("named",)


def load_profile_fmte02(path: str) -> dict:
    """e-v0.2 읽기 설정을 읽고 칸을 검사한다. 모르는 칸·빠진 칸·잘못된 값이 있으면 멈춘다."""
    import yaml

    with open(path, encoding="utf-8") as f:
        raw = yaml.safe_load(f)
    if not isinstance(raw, dict):
        raise XlsxGoldError(f"{path}: 설정은 키와 값의 목록이어야 합니다")
    extra, missing = set(raw) - _NET_KEYS, _NET_KEYS - set(raw)
    if extra or missing:
        raise XlsxGoldError(f"{path}: 칸이 맞지 않습니다(모르는 칸 {sorted(extra)}, 빠진 칸 {sorted(missing)})")
    if raw["profile"] != "e-v0.2":
        raise XlsxGoldError(f"{path}: profile은 'e-v0.2'이어야 합니다: {raw['profile']!r}")
    if raw["bundle_mode"] not in _NET_BUNDLE_MODES:
        raise XlsxGoldError(f"{path}: bundle_mode는 {list(_NET_BUNDLE_MODES)} 중 하나여야 합니다: {raw['bundle_mode']!r}")
    fmtb_like = dict(raw, bundle_mode="blank_single", question_subset="all", issue_questions={"ids": ["(없음)"], "reason": "(없음)"})
    fmtb_like["profile"] = "b-v0.1"
    import tempfile
    with tempfile.NamedTemporaryFile("w", suffix=".yaml", delete=False, encoding="utf-8") as tf:      # 나머지 칸 검사는 fmtb01과 같다
        yaml.safe_dump(fmtb_like, tf, allow_unicode=True)
        tmp = tf.name
    try:
        load_profile_fmtb01(tmp)
    finally:
        os.remove(tmp)
    return raw


def read_gold_profile_fmte02(path: str, profile_path: str):
    """서비스 E 평가셋 v0.2를 설정대로 읽는다. (정답표 목록, 변환 보고)를 돌려주고 어긋나면 XlsxGoldError."""
    import openpyxl

    prof = load_profile_fmte02(profile_path)
    fmt = scoring.load_format()
    wb = openpyxl.load_workbook(path, data_only=True)
    sh = prof["sheets"]
    for s_ in sh.values():
        if s_ not in wb.sheetnames:
            raise XlsxGoldError(f"시트가 없습니다: {s_}")
    hdr, first = prof["header_row"], prof["first_row"]
    qi, q_unused = _profile_header(wb[sh["questions"]], hdr, prof["question_columns"], sh["questions"])
    li, l_unused = _profile_header(wb[sh["question_list"]], hdr, prof["question_list_columns"], sh["question_list"])
    ei, e_unused = _profile_header(wb[sh["evidence"]], hdr, prof["evidence_columns"], sh["evidence"])
    di, d_unused = _profile_header(wb[sh["documents"]], hdr, prof["document_columns"], sh["documents"])

    # 기준 문서 시트: 파일명 -> 문서 ID, 지문 대조
    sheet_docs, file_to_doc = {}, {}
    for r in wb[sh["documents"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        if _blank(c[di["id"]]):
            continue
        did = str(c[di["id"]]).strip()
        fname = _file_name(str(c[di["file"]]), f"{sh['documents']} {r[0].row}행")
        if did in sheet_docs or fname in file_to_doc:
            raise XlsxGoldError(f"{sh['documents']} {r[0].row}행: 문서 ID 또는 파일명이 중복입니다")
        file_to_doc[fname] = did
        sheet_docs[did] = {"file": fname, "sha256": None if _blank(c[di["sha256"]]) else str(c[di["sha256"]]).strip().lower(),
                           "state": None if _blank(c[di["state"]]) else str(c[di["state"]]).strip()}
    doc_status = {}
    for did, spec in prof["documents"].items():
        got = sheet_docs.get(did)
        if got is None:
            raise XlsxGoldError(f"문서 {did}: 평가셋 기준 문서 시트에 없습니다")
        if got["file"] != spec["file"]:
            raise XlsxGoldError(f"문서 {did}: 설정의 파일 이름 {spec['file']!r}이 기준 문서 시트의 {got['file']!r}과 다릅니다")
        if got["sha256"] != spec["sha256"].lower():
            raise XlsxGoldError(f"문서 {did}: 지문이 다릅니다(평가셋 {str(got['sha256'])[:12]} 대 설정 {spec['sha256'][:12]})")
        doc_status[did] = {"status": "ok", "state": got["state"], "reason": None, "hash_mismatch": False}

    # 질문 정답표와 질문 목록 대조
    questions, by_id, states = [], {}, {}
    for r in wb[sh["questions"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        if _blank(c[qi["id"]]):
            continue
        qid = str(c[qi["id"]]).strip()
        where = f"{sh['questions']} {r[0].row}행({qid})"
        if qid in by_id:
            raise XlsxGoldError(f"{where}: 질문 ID 중복")
        src_type = _text(c[qi["type"]], where, "유형", required=True)
        if src_type not in prof["type_map"] or prof["type_map"][src_type] not in fmt["types"]:
            raise XlsxGoldError(f"{where}: 유형 {src_type!r}을 읽을 수 없습니다(type_map 또는 gold_format.yaml types에 없음)")
        sv = "(빈칸)" if _blank(c[qi["state"]]) else str(c[qi["state"]]).strip()
        states[sv] = states.get(sv, 0) + 1
        d = {"case_id": qid, "type": prof["type_map"][src_type], "request": _text(c[qi["text"]], where, "질문", required=True),
             "gold_evidence": [], "slot_rule": scoring.UNIT_BUNDLE}
        by_id[qid] = {"d": d, "units": {}}
        questions.append(d)
    if not questions:
        raise XlsxGoldError(f"{sh['questions']}에 읽을 질문이 하나도 없습니다")
    listed = {}
    for r in wb[sh["question_list"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        if not _blank(c[li["id"]]):
            listed[str(c[li["id"]]).strip()] = None if _blank(c[li["text"]]) else str(c[li["text"]]).strip()
    if set(listed) != set(by_id) or any(listed[q] != by_id[q]["d"]["request"] for q in by_id):
        raise XlsxGoldError("질문 목록과 질문 정답표의 질문 ID 또는 질문 글이 다릅니다")

    # 질문별 근거: 대체 묶음 값을 그대로 묶음 이름으로(빈칸이면 멈춤)
    n_evidence = 0
    bundles_seen = {}
    for r in wb[sh["evidence"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        if _blank(c[ei["question_id"]]):
            if any(not _blank(c[ei[k]]) for k in ("unit", "bundle", "role", "file", "sheet_cells")):
                raise XlsxGoldError(f"{sh['evidence']} {r[0].row}행: 질문 ID가 비었는데 읽는 열에 값이 있습니다")
            continue
        qid = str(c[ei["question_id"]]).strip()
        where = f"{sh['evidence']} {r[0].row}행({qid})"
        if qid not in by_id:
            raise XlsxGoldError(f"{where}: 질문 정답표에 없는 질문 ID")
        unit = _text(c[ei["unit"]], where, "채점 단위 ID", required=True)
        bundle = _text(c[ei["bundle"]], where, "대체 묶음", required=True)
        bundles_seen[bundle] = bundles_seen.get(bundle, 0) + 1
        role = _text(c[ei["role"]], where, "필수/보조", required=True)
        if role not in prof["role_values"]:
            raise XlsxGoldError(f"{where}: 필수/보조 값 {role!r}을 모릅니다")
        fname = _file_name(_text(c[ei["file"]], where, "파일명", required=True), where)
        if fname not in file_to_doc or file_to_doc[fname] not in prof["documents"]:
            raise XlsxGoldError(f"{where}: 파일명 {fname!r}이 기준 문서 또는 설정의 documents에 없습니다")
        sc = c[ei["sheet_cells"]]
        if not isinstance(sc, str) or "!" not in sc:
            raise XlsxGoldError(f"{where}: 시트·셀 칸이 '시트!셀범위' 모양이 아닙니다: {sc!r}")
        sheet, cells = sc.rsplit("!", 1)                              # 마지막 !에서 나눈다. 시트 이름은 글자 하나도 바꾸지 않는다
        if not sheet.strip():
            raise XlsxGoldError(f"{where}: 시트 이름이 비었습니다")
        ev = {"file": fname, "sheet": sheet, "cells": _cell_list(cells, where), "required": bool(prof["role_values"][role]),
              "unit": unit, "bundle": bundle}
        by_id[qid]["d"]["gold_evidence"].append(ev)
        by_id[qid]["units"].setdefault(unit, []).append((bundle, ev["required"]))
        n_evidence += 1

    req_rows = req_units = opt_units = alt_units = 0
    for qid, q in by_id.items():
        for unit, rows in q["units"].items():
            if len({x[1] for x in rows}) > 1:
                raise XlsxGoldError(f"질문 {qid}: 채점 단위 {unit!r} 안에 필수와 보조가 섞여 있습니다")
            if rows[0][1]:
                req_units += 1
                req_rows += len(rows)
                alt_units += len({x[0] for x in rows}) > 1
            else:
                opt_units += 1
    got_counts = {"questions": len(questions), "evidence": n_evidence, "required_rows": req_rows, "required_units": req_units}
    bad = {k: (got_counts[k], prof["expected_counts"][k]) for k in _V02_COUNTS if got_counts[k] != prof["expected_counts"][k]}
    if bad:
        raise XlsxGoldError("읽은 개수가 설정의 expected_counts와 다릅니다(읽은 값, 기대 값): " + json.dumps(bad, ensure_ascii=False))
    try:
        scoring.parse_gold(questions)
    except scoring.ScoringError as err:
        raise XlsxGoldError(f"채점기 검사에서 멈춤: {err}") from None
    by_type = dict(sorted({t: sum(1 for d in questions if d["type"] == t) for t in {d["type"] for d in questions}}.items()))
    profile_report = {"name": prof["profile"], "rule": prof["rule"], "documents": doc_status, "held_questions": [],
                      "n_questions_before_hold": len(questions), "n_evidence_before_hold": n_evidence,
                      "units_before_hold": {"required": req_units, "optional": opt_units}, "required_rows": req_rows,
                      "by_type_before_hold": by_type, "by_type_after_hold": by_type,
                      "bundle_mode": {"mode": prof["bundle_mode"], "values": dict(sorted(bundles_seen.items())),
                                      "required_units_with_alternatives": alt_units,
                                      "meaning": "대체 묶음 값을 묶음 이름으로 씀. 같은 묶음 AND, 다른 묶음 OR"},
                      "question_subset": {"subset": "all", "issue_ids": [], "reason": "쟁점 문항 없음", "left_out": [], "before_subset": got_counts},
                      "not_available": {"답변 가능 범위": "열 없음, 채점기 기본값(가능)을 씀", "적용 조건": "열 없음, 조건부 채점을 하지 않음",
                                        "검토 영역": "열 없음, 영역별 누락 보고 없음"},
                      "source_state": {"상태": states, "문서 출처 상태": {k: v["state"] for k, v in doc_status.items()}}}
    return questions, {"unused_columns": {sh["questions"]: sorted(q_unused), sh["question_list"]: sorted(l_unused),
                                          sh["evidence"]: sorted(e_unused), sh["documents"]: sorted(d_unused)},
                       "profile": profile_report}


# ---------- 묶음 82: 서비스 C 평가셋 읽기(기준은 CHANGES.md 묶음 82 S1, 서비스 B 함수를 복사해 위치 칸만 바꿈) ----------
_FMTC_ECOLS = ("question_id", "unit", "bundle", "role", "file", "sheet_rows")


def load_profile_fmtc01(path: str) -> dict:
    """묶음 82 S1: 서비스 C 평가셋 읽기 설정. load_profile_fmtb01과 같고 profile 이름과 근거 위치 칸(sheet_rows)만 다르다."""
    import yaml

    with open(path, encoding="utf-8") as f:
        raw = yaml.safe_load(f)
    if not isinstance(raw, dict):
        raise XlsxGoldError(f"{path}: 설정은 키와 값의 목록이어야 합니다")
    extra, missing = set(raw) - _FMTB_KEYS, _FMTB_KEYS - set(raw)
    if extra or missing:
        raise XlsxGoldError(f"{path}: 칸이 맞지 않습니다(모르는 칸 {sorted(extra)}, 빠진 칸 {sorted(missing)})")
    if raw["profile"] != "c-v0.1":
        raise XlsxGoldError(f"{path}: profile은 'c-v0.1'이어야 합니다: {raw['profile']!r}")
    if raw["rule"] != scoring.UNIT_BUNDLE:
        raise XlsxGoldError(f"{path}: rule은 {scoring.UNIT_BUNDLE!r}만 쓸 수 있습니다: {raw['rule']!r}")
    for key, names in (("question_columns", _V02_QCOLS), ("question_list_columns", _V02_LCOLS), ("evidence_columns", _FMTC_ECOLS),
                       ("document_columns", _V02_DCOLS)):
        got = raw[key]
        if not isinstance(got, dict) or set(got) != set(names) or not all(isinstance(v, str) and v.strip() for v in got.values()):
            raise XlsxGoldError(f"{path}: {key}는 {list(names)}를 모두 가진 비어 있지 않은 문자열이어야 합니다")
    if not isinstance(raw["sheets"], dict) or set(raw["sheets"]) != {"questions", "question_list", "evidence", "documents"} \
            or not all(isinstance(v, str) and v for v in raw["sheets"].values()):
        raise XlsxGoldError(f"{path}: sheets는 questions, question_list, evidence, documents 이름을 가져야 합니다")
    for k in ("header_row", "first_row"):
        if isinstance(raw[k], bool) or not isinstance(raw[k], int) or raw[k] < 1:
            raise XlsxGoldError(f"{path}: {k}는 1 이상의 정수여야 합니다")
    if not isinstance(raw["role_values"], dict) or not raw["role_values"] or not all(isinstance(b, bool) for b in raw["role_values"].values()):
        raise XlsxGoldError(f"{path}: role_values는 값 이름 -> true/false 목록이어야 합니다")
    tm = raw["type_map"]
    if not isinstance(tm, dict) or not tm or not all(isinstance(k, str) and isinstance(v, str) and k.strip() and v.strip() for k, v in tm.items()):
        raise XlsxGoldError(f"{path}: type_map은 원본 유형 이름 -> 채점기 유형 이름 목록이어야 합니다")
    ec = raw["expected_counts"]
    if not isinstance(ec, dict) or set(ec) != set(_V02_COUNTS) or not all(isinstance(v, int) and not isinstance(v, bool) and v >= 0 for v in ec.values()):
        raise XlsxGoldError(f"{path}: expected_counts는 {list(_V02_COUNTS)}를 모두 가진 0 이상의 정수여야 합니다")
    if not isinstance(raw["documents"], dict) or not raw["documents"]:
        raise XlsxGoldError(f"{path}: documents가 비어 있습니다")
    for doc, v in raw["documents"].items():
        if not isinstance(v, dict) or set(v) != {"file", "sha256"} or not all(isinstance(v[k], str) and v[k].strip() for k in ("file", "sha256")):
            raise XlsxGoldError(f"{path}: documents.{doc}는 file과 sha256만 가져야 합니다")
        _file_name(v["file"], f"{path} documents.{doc}")
    if raw["bundle_mode"] not in _FMTB_BUNDLE_MODES:
        raise XlsxGoldError(f"{path}: bundle_mode는 {list(_FMTB_BUNDLE_MODES)} 중 하나여야 합니다: {raw['bundle_mode']!r}")
    if raw["question_subset"] not in _FMTB_SUBSETS:
        raise XlsxGoldError(f"{path}: question_subset은 {list(_FMTB_SUBSETS)} 중 하나여야 합니다: {raw['question_subset']!r}")
    iq = raw["issue_questions"]
    if not isinstance(iq, dict) or set(iq) != {"ids", "reason"}:
        raise XlsxGoldError(f"{path}: issue_questions는 ids와 reason만 가져야 합니다")
    ids = iq["ids"]
    if not isinstance(ids, list) or not ids or not all(isinstance(x, str) and x.strip() for x in ids):
        raise XlsxGoldError(f"{path}: issue_questions.ids는 비어 있지 않은 질문 ID 목록이어야 합니다")
    dup = sorted({x for x in ids if ids.count(x) > 1})
    if dup:
        raise XlsxGoldError(f"{path}: issue_questions.ids에 같은 ID가 두 번 이상 있습니다: {dup}")
    if not isinstance(iq["reason"], str) or not iq["reason"].strip():
        raise XlsxGoldError(f"{path}: issue_questions.reason이 비었습니다")
    return raw


def read_gold_profile_fmtc01(path: str, profile_path: str):
    """묶음 82 S1: 서비스 C 평가셋을 설정대로 읽는다. read_gold_profile_fmtb01과 같고 근거 위치가 "시트!N행"(행 하나)이다."""
    import openpyxl

    prof = load_profile_fmtc01(profile_path)
    fmt = scoring.load_format()
    wb = openpyxl.load_workbook(path, data_only=True)
    sh = prof["sheets"]
    for s_ in sh.values():
        if s_ not in wb.sheetnames:
            raise XlsxGoldError(f"시트가 없습니다: {s_}")
    hdr, first = prof["header_row"], prof["first_row"]
    qi, q_unused = _profile_header(wb[sh["questions"]], hdr, prof["question_columns"], sh["questions"])
    li, l_unused = _profile_header(wb[sh["question_list"]], hdr, prof["question_list_columns"], sh["question_list"])
    ei, e_unused = _profile_header(wb[sh["evidence"]], hdr, prof["evidence_columns"], sh["evidence"])
    di, d_unused = _profile_header(wb[sh["documents"]], hdr, prof["document_columns"], sh["documents"])

    # 기준 문서 시트(파일명 -> 문서 ID)와 지문. 상태는 코드가 계산한다(v0.2와 같음, 풀어 주기 칸은 없음)
    sheet_docs, file_to_doc = {}, {}
    for r in wb[sh["documents"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        did = c[di["id"]]
        if _blank(did):
            continue
        did = str(did).strip()
        fname = c[di["file"]]
        if _blank(fname) or did in sheet_docs:
            raise XlsxGoldError(f"{sh['documents']} {r[0].row}행: 파일명이 비었거나 문서 ID({did})가 중복입니다")
        fname = _file_name(str(fname), f"{sh['documents']} {r[0].row}행")
        if fname in file_to_doc:
            raise XlsxGoldError(f"{sh['documents']} {r[0].row}행: 파일명 {fname!r}이 둘 이상의 문서 ID에 있습니다")
        file_to_doc[fname] = did
        sheet_docs[did] = {"file": fname, "sha256": None if _blank(c[di["sha256"]]) else str(c[di["sha256"]]).strip().lower(),
                           "state": None if _blank(c[di["state"]]) else str(c[di["state"]]).strip()}
    doc_status = {}
    for did, spec in prof["documents"].items():
        got = sheet_docs.get(did)
        if got is None:
            doc_status[did] = {"status": "unverifiable", "state": None, "reason": "평가셋 기준 문서 시트에 없음"}
        elif got["file"] != spec["file"]:
            raise XlsxGoldError(f"문서 {did}: 설정의 파일 이름 {spec['file']!r}이 평가셋 기준 문서 시트의 {got['file']!r}과 다릅니다")
        elif got["sha256"] == spec["sha256"].lower():
            doc_status[did] = {"status": "ok", "state": got["state"], "reason": None, "hash_mismatch": False}
        else:
            doc_status[did] = {"status": "hold", "state": got["state"], "hash_mismatch": True,
                               "reason": f"문서 지문 불일치(평가셋 {str(got['sha256'])[:12]} 대 우리 {spec['sha256'][:12]})"}

    # 질문 정답표(유형은 type_map으로 바꾸고 채점기 목록에 있어야 한다)
    questions, by_id, states = [], {}, {}
    for r in wb[sh["questions"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        qid = c[qi["id"]]
        where = f"{sh['questions']} {r[0].row}행"
        if _blank(qid):
            if any(not _blank(c[qi[k]]) for k in ("type", "text")):
                raise XlsxGoldError(f"{where}: 질문 ID가 비었는데 읽는 열에 값이 있습니다")
            continue
        if not isinstance(qid, str):
            raise XlsxGoldError(f"{where}: 질문 ID가 글이 아닙니다")
        qid = qid.strip()
        where = f"{where}({qid})"
        if qid in by_id:
            raise XlsxGoldError(f"{where}: 질문 ID 중복")
        src_type = _text(c[qi["type"]], where, "유형", required=True)
        if src_type not in prof["type_map"]:
            raise XlsxGoldError(f"{where}: 유형 {src_type!r}이 설정의 type_map에 없습니다(있는 이름 {sorted(prof['type_map'])})")
        qtype = prof["type_map"][src_type]
        if qtype not in fmt["types"]:
            raise XlsxGoldError(f"{where}: 유형 {qtype!r}이 gold_format.yaml의 types에 없습니다(채점기가 읽지 못함)")
        text = _text(c[qi["text"]], where, "질문", required=True)
        sv = c[qi["state"]]
        sv = "(빈칸)" if _blank(sv) else str(sv).strip()
        states[sv] = states.get(sv, 0) + 1
        d = {"case_id": qid, "type": qtype, "request": text, "gold_evidence": [], "slot_rule": scoring.UNIT_BUNDLE}
        by_id[qid] = {"d": d, "units": {}}
        questions.append(d)
    if not questions:
        raise XlsxGoldError(f"{sh['questions']}에 읽을 질문이 하나도 없습니다")

    # 질문 목록 시트와 같은지 대조
    listed = {}
    for r in wb[sh["question_list"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        if _blank(c[li["id"]]):
            continue
        qid = str(c[li["id"]]).strip()
        if qid in listed:
            raise XlsxGoldError(f"{sh['question_list']} {r[0].row}행: 질문 ID {qid} 중복")
        listed[qid] = None if _blank(c[li["text"]]) else str(c[li["text"]]).strip()
    if set(listed) != set(by_id):
        raise XlsxGoldError(f"질문 목록과 질문 정답표의 질문 ID가 다릅니다: 목록에만 {sorted(set(listed) - set(by_id))}, "
                            f"정답표에만 {sorted(set(by_id) - set(listed))}")
    diff_text = sorted(q for q in by_id if listed[q] != by_id[q]["d"]["request"])
    if diff_text:
        raise XlsxGoldError(f"질문 목록과 질문 정답표의 질문 글이 다른 질문이 있습니다: {diff_text}")

    # 질문별 근거(D1a: 대체 묶음은 모두 빈칸이어야 하고 단위마다 묶음 하나)
    n_evidence, blank_rows = 0, 0
    for r in wb[sh["evidence"]].iter_rows(min_row=first):
        c = [x.value for x in r]
        qid = c[ei["question_id"]]
        read = [c[ei[k]] for k in ("unit", "bundle", "role", "file", "sheet_rows")]
        if _blank(qid):
            if any(not _blank(v) for v in read):
                raise XlsxGoldError(f"{sh['evidence']} {r[0].row}행: 질문 ID가 비었는데 읽는 열에 값이 있습니다")
            blank_rows += 1
            continue
        qid = str(qid).strip()
        where = f"{sh['evidence']} {r[0].row}행({qid})"
        if qid not in by_id:
            raise XlsxGoldError(f"{where}: 질문 정답표에 없는 질문 ID")
        unit = _text(c[ei["unit"]], where, "채점 단위 ID", required=True)
        if not _blank(c[ei["bundle"]]):
            raise XlsxGoldError(f"{where}: bundle_mode {prof['bundle_mode']!r}인데 대체 묶음 칸에 값이 있습니다: {c[ei['bundle']]!r}")
        bundle = FMTB_SINGLE_BUNDLE
        role = _text(c[ei["role"]], where, "필수/보조", required=True)
        if role not in prof["role_values"]:
            raise XlsxGoldError(f"{where}: 필수/보조 값 {role!r}을 모릅니다(설정에 있는 값 {list(prof['role_values'])})")
        fname = _file_name(_text(c[ei["file"]], where, "파일명", required=True), where)
        if fname not in file_to_doc:
            raise XlsxGoldError(f"{where}: 파일명 {fname!r}이 평가셋 '{sh['documents']}' 시트에 없습니다")
        did = file_to_doc[fname]
        if did not in prof["documents"]:
            raise XlsxGoldError(f"{where}: 문서 ID {did!r}가 설정의 documents에 없습니다")
        if doc_status[did]["status"] == "unverifiable":
            raise XlsxGoldError(f"{where}: 문서 {did}의 지문을 평가셋 '{sh['documents']}' 시트에서 확인할 수 없습니다")
        sc = c[ei["sheet_rows"]]
        if not isinstance(sc, str) or "!" not in sc:
            raise XlsxGoldError(f"{where}: 시트·행 칸이 '시트!N행' 모양이 아닙니다: {sc!r}")
        sheet, rowpart = sc.rsplit("!", 1)                            # 마지막 !에서 나눈다. 시트 이름은 글자 하나도 바꾸지 않는다(이슈 26)
        if not sheet.strip():
            raise XlsxGoldError(f"{where}: 시트 이름이 비었습니다: {sc!r}")
        m = re.fullmatch(r"([1-9][0-9]*)행", rowpart)
        if not m:
            raise XlsxGoldError(f"{where}: 시트·행 칸의 행이 'N행' 모양이 아닙니다: {sc!r}")
        n_row = int(m.group(1))
        ev = {"file": fname, "sheet": sheet, "rows": [n_row, n_row], "required": bool(prof["role_values"][role]), "unit": unit, "bundle": bundle}
        q = by_id[qid]
        q["d"]["gold_evidence"].append(ev)
        q["units"].setdefault(unit, []).append((bundle, ev["required"], did))
        n_evidence += 1

    def unit_counts(ids):
        rr = ru = ou = 0
        for qid in ids:
            for unit, rows in by_id[qid]["units"].items():
                if rows[0][1]:
                    ru += 1
                    rr += len(rows)
                else:
                    ou += 1
        return rr, ru, ou

    for qid, q in by_id.items():
        for unit, rows in q["units"].items():
            if len({x[1] for x in rows}) > 1:
                raise XlsxGoldError(f"질문 {qid}: 채점 단위 {unit!r} 안에 필수와 보조가 섞여 있습니다")
    req_rows, req_units, opt_units = unit_counts(by_id)
    shapes = {"단일 행": 0, "여러 행": 0}
    for q in by_id.values():
        for rows in q["units"].values():
            shapes["단일 행" if len(rows) == 1 else "여러 행"] += 1
    # D1d: 기대 개수는 빼기 전 개수로 대조한다
    got_counts = {"questions": len(questions), "evidence": n_evidence, "required_rows": req_rows, "required_units": req_units}
    bad = {k: (got_counts[k], prof["expected_counts"][k]) for k in _V02_COUNTS if got_counts[k] != prof["expected_counts"][k]}
    if bad:
        raise XlsxGoldError("읽은 개수가 설정의 expected_counts와 다릅니다(읽은 값, 기대 값): " + json.dumps(bad, ensure_ascii=False))

    # D1c: 쟁점 질문을 빼거나(main) 그것만 남긴다(issues)
    issue_ids = [x.strip() for x in prof["issue_questions"]["ids"]]
    unknown = sorted(set(issue_ids) - set(by_id))
    if unknown:
        raise XlsxGoldError(f"issue_questions.ids에 정답지에 없는 질문 ID가 있습니다: {unknown}")
    subset = prof["question_subset"]
    chosen = list(questions) if subset == "all" else [d for d in questions if (d["case_id"] in issue_ids) == (subset == "issues")]
    if not chosen:
        raise XlsxGoldError(f"question_subset {subset!r}로 고르면 남는 질문이 없습니다")
    left_out = [d["case_id"] for d in questions if d not in chosen]

    held, kept = [], []
    for d in chosen:
        docs = {}
        for rows in by_id[d["case_id"]]["units"].values():
            for _b, role_, did in rows:
                if doc_status[did]["status"] == "hold":
                    docs.setdefault(did, set()).add("필수" if role_ else "보조")
        if docs:
            held.append({"id": d["case_id"], "documents": sorted(docs), "evidence_roles": {k: sorted(v) for k, v in sorted(docs.items())},
                         "reason": "; ".join(f"{k}: {doc_status[k]['reason']}" for k in sorted(docs))})
        else:
            kept.append(d)
    if not kept:
        raise XlsxGoldError("보류를 빼면 남는 질문이 없습니다")
    try:
        scoring.parse_gold(kept)
    except scoring.ScoringError as err:
        raise XlsxGoldError(f"채점기 검사에서 멈춤: {err}") from None
    sub_rr, sub_ru, sub_ou = unit_counts([d["case_id"] for d in chosen])
    by_type = lambda qs: dict(sorted({t: sum(1 for d in qs if d["type"] == t) for t in {d["type"] for d in qs}}.items()))
    profile_report = {"name": prof["profile"], "rule": prof["rule"], "documents": doc_status, "held_questions": held,
                      "n_questions_before_hold": len(chosen), "n_evidence_before_hold": sum(len(d["gold_evidence"]) for d in chosen),
                      "units_before_hold": {"required": sub_ru, "optional": sub_ou}, "unit_shapes": shapes,
                      "row_kinds": {}, "blank_evidence_rows": blank_rows, "required_rows": sub_rr,
                      "by_type_before_hold": by_type(chosen), "by_type_after_hold": by_type(kept),
                      "bundle_mode": {"mode": prof["bundle_mode"], "read_as": FMTB_SINGLE_BUNDLE,
                                      "meaning": "대체 묶음이 모두 빈칸이라 단위마다 묶음 하나(단위 안 근거는 모두 필요)"},
                      "question_subset": {"subset": subset, "issue_ids": issue_ids, "reason": prof["issue_questions"]["reason"].strip(),
                                          "left_out": left_out, "before_subset": got_counts},
                      "not_available": {"답변 가능 범위": "열 없음, 채점기 기본값(가능)을 씀", "적용 조건": "열 없음, 조건부 채점을 하지 않음",
                                        "검토 영역": "열 없음, 영역별 누락 보고 없음"},
                      "source_state": {"상태": states, "문서 출처 상태": {k: v["state"] for k, v in doc_status.items()}}}
    return kept, {"unused_columns": {sh["questions"]: sorted(q_unused), sh["question_list"]: sorted(l_unused), sh["evidence"]: sorted(e_unused),
                                     sh["documents"]: sorted(d_unused)},
                  "profile": profile_report}


# ---------- 묶음 97: 서비스 D 평가셋 읽기(기준은 CHANGES.md 묶음 97 W7) ----------
# 서비스 C와 같은 양식이라 서비스 C 함수를 그대로 부른다(기존 함수는 고치지 않음). 다른 점은 쟁점 문항이 없다는 것 하나:
#   설정에 issue_questions·question_subset 칸이 없고 30문항을 모두 읽는다.
#   서비스 C 함수는 쟁점 ID가 정답지에 있어야 하므로, 임시 설정에 첫 질문 ID를 넣고 subset "all"로 부른 뒤(아무것도 빼지 않음)
#   보고의 쟁점 칸을 "쟁점 문항 없음"으로 바로잡는다. 문서 지문이 달라 보류되는 질문이 있으면 멈춘다(서비스 E와 같게).
_FMTD_KEYS = _FMTB_KEYS - {"issue_questions", "question_subset"}


def _fmtc_like(raw: dict, issue_id: str) -> dict:
    return dict(raw, profile="c-v0.1", question_subset="all", issue_questions={"ids": [issue_id], "reason": "(쟁점 문항 없음)"})


def _dump_tmp_yaml(d: dict) -> str:
    import tempfile
    import yaml

    with tempfile.NamedTemporaryFile("w", suffix=".yaml", delete=False, encoding="utf-8") as tf:
        yaml.safe_dump(d, tf, allow_unicode=True)
        return tf.name


def load_profile_fmtd01(path: str) -> dict:
    """d-v0.1 읽기 설정을 읽고 칸을 검사한다. 나머지 칸 검사는 서비스 C와 같다. 어긋나면 멈춘다."""
    import yaml

    with open(path, encoding="utf-8") as f:
        raw = yaml.safe_load(f)
    if not isinstance(raw, dict):
        raise XlsxGoldError(f"{path}: 설정은 키와 값의 목록이어야 합니다")
    extra, missing = set(raw) - _FMTD_KEYS, _FMTD_KEYS - set(raw)
    if extra or missing:
        raise XlsxGoldError(f"{path}: 칸이 맞지 않습니다(모르는 칸 {sorted(extra)}, 빠진 칸 {sorted(missing)})")
    if raw["profile"] != "d-v0.1":
        raise XlsxGoldError(f"{path}: profile은 'd-v0.1'이어야 합니다: {raw['profile']!r}")
    if raw["bundle_mode"] != "blank_single":
        raise XlsxGoldError(f"{path}: bundle_mode는 'blank_single'만 쓸 수 있습니다: {raw['bundle_mode']!r}")
    tmp = _dump_tmp_yaml(_fmtc_like(raw, "(없음)"))
    try:
        load_profile_fmtc01(tmp)
    finally:
        os.remove(tmp)
    return raw


def read_gold_profile_fmtd01(path: str, profile_path: str):
    """서비스 D 평가셋을 설정대로 읽는다(30문항 전부). (정답표 목록, 변환 보고)를 돌려주고 어긋나면 XlsxGoldError."""
    import openpyxl

    prof = load_profile_fmtd01(profile_path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sh = prof["sheets"]["questions"]
    if sh not in wb.sheetnames:
        raise XlsxGoldError(f"시트가 없습니다: {sh}")
    qi, _ = _profile_header(wb[sh], prof["header_row"], prof["question_columns"], sh)
    first_id = next((str(r[qi["id"]]).strip() for r in wb[sh].iter_rows(min_row=prof["first_row"], values_only=True)
                     if not _blank(r[qi["id"]])), None)
    wb.close()
    if first_id is None:
        raise XlsxGoldError(f"{sh}에 읽을 질문이 하나도 없습니다")
    tmp = _dump_tmp_yaml(_fmtc_like(prof, first_id))
    try:
        questions, report = read_gold_profile_fmtc01(path, tmp)
    finally:
        os.remove(tmp)
    p = report["profile"]
    if p["held_questions"]:
        raise XlsxGoldError("문서 지문이 달라 보류된 질문이 있습니다(멈춤): " + json.dumps([h["id"] for h in p["held_questions"]], ensure_ascii=False))
    if p["question_subset"]["left_out"]:
        raise XlsxGoldError(f"빠진 질문이 있습니다(30문항 전부여야 함): {p['question_subset']['left_out']}")
    p["name"] = prof["profile"]
    p["question_subset"] = {"subset": "all", "issue_ids": [], "reason": "쟁점 문항 없음", "left_out": [],
                            "before_subset": p["question_subset"]["before_subset"]}
    return questions, report


# ---------- 스냅샷과 변경 감지 (N6~N8) ----------
def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _canon(obj) -> bytes:
    return json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def _json_bytes(questions) -> bytes:
    return json.dumps(questions, ensure_ascii=False, indent=1).encode("utf-8") + b"\n"


def _file_sha(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def hidden_rows_report(xlsx_path: str, sheets: list, first_row: int) -> dict:
    """묶음 ㉞ A1: 읽는 범위(첫 데이터 행부터)에서 숨긴 행 가운데 값이 있는 행의 번호를 시트별로 센다. 값(본문)은 담지 않는다.
    숨긴 행을 삭제로 볼지는 코드가 정하지 않고, 읽는 범위도 바꾸지 않는다."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    out = {}
    for name in sheets:
        ws = wb[name]
        rows = [r for r in range(first_row, ws.max_row + 1)
                if ws.row_dimensions[r].hidden and any(not _blank(c.value) for c in ws[r])]
        out[name] = rows
    return out


def condition_report(questions: list) -> dict:
    """묶음 ㉞ A2: 조건 값별 근거 수와, 질문에 명시되지 않았고 '항상'도 아닌 조건 값(질문 id별)을 보고한다. 판단하지 않는다."""
    always = scoring.load_format()["always_word"]
    counts, unstated = {}, {}
    for q in questions:
        stated = set(q.get("stated_conditions") or [])
        for e in q["gold_evidence"]:
            c = e.get("condition")
            if c is None:
                continue
            counts[c] = counts.get(c, 0) + 1
            if c != always and c not in stated:
                unstated.setdefault(q["case_id"], [])
                if c not in unstated[q["case_id"]]:
                    unstated[q["case_id"]].append(c)
    return {"counts": dict(sorted(counts.items())), "unstated": unstated}


def build_snapshot(xlsx_path: str, questions: list, profile_info: dict | None = None, unused_columns: dict | None = None,
                   hidden_rows: dict | None = None, condition_counts: dict | None = None) -> dict:
    by_type = {}
    for q in questions:
        by_type[q["type"]] = by_type.get(q["type"], 0) + 1
    per = {}
    for q in questions:
        ev = q["gold_evidence"]
        per[q["case_id"]] = {"required": sum(1 for e in ev if e["required"]), "optional": sum(1 for e in ev if not e["required"]),
                             "digest": _sha(_canon(q))}
    snap = {"snapshot_version": 1, "gold_file": os.path.basename(xlsx_path), "gold_sha256": _file_sha(xlsx_path),
            "questions": len(questions), "evidence": sum(len(q["gold_evidence"]) for q in questions),
            "by_type": dict(sorted(by_type.items())), "per_question": per, "json_sha256": _sha(_json_bytes(questions)),
            "status": UNCONFIRMED, "confirmed_by": None, "confirmed_date": None, "reason": None}
    if profile_info is not None:                                    # 묶음 ㉒ V23: 읽기 설정을 쓴 정답표만 이 칸이 있다(옛 스냅샷은 그대로)
        snap["profile"] = profile_info
    if unused_columns is not None:                                  # 묶음 ㉝ Y4: 읽지 않은 열을 남겨 두면 새 열이 생겼을 때 알 수 있다
        snap["unused_columns"] = {k: sorted(v) for k, v in unused_columns.items()}
    if hidden_rows is not None:                                     # 묶음 ㉞ A1: 행 번호만 저장한다(값은 저장하지 않음)
        snap["hidden_rows"] = {k: list(v) for k, v in hidden_rows.items()}
    if condition_counts is not None:                                # 묶음 ㉞ A2
        snap["condition_counts"] = dict(condition_counts)
    return snap


def diff_snapshots(old: dict, new: dict) -> list:
    """달라진 점을 사람이 읽을 문장 목록으로. 같으면 []."""
    out = []
    if old["gold_file"] != new["gold_file"]:
        out.append(f"정답지 파일 이름: {old['gold_file']} -> {new['gold_file']}")
    if old["gold_sha256"] != new["gold_sha256"]:
        out.append(f"정답지 파일 지문이 다릅니다({old['gold_sha256'][:12]} -> {new['gold_sha256'][:12]})")
    for k, label in (("questions", "질문 수"), ("evidence", "근거 수")):
        if old[k] != new[k]:
            out.append(f"{label}: {old[k]} -> {new[k]}")
    if old["by_type"] != new["by_type"]:
        for t in sorted(set(old["by_type"]) | set(new["by_type"])):
            a, b = old["by_type"].get(t, 0), new["by_type"].get(t, 0)
            if a != b:
                out.append(f"유형 '{t}' 질문 수: {a} -> {b}")
    uo, un = old.get("unused_columns"), new.get("unused_columns")
    if uo is not None and un is not None:                           # 묶음 ㉝ Y4: 칸이 없는 옛 스냅샷은 건너뛴다
        for sheet in sorted(set(uo) | set(un)):
            added = sorted(set(un.get(sheet, [])) - set(uo.get(sheet, [])))
            if added:
                out.append(f"시트 '{sheet}'에 쓰지 않는 새 열이 생겼습니다: {added}")
    ho, hn = old.get("hidden_rows"), new.get("hidden_rows")
    if ho is not None and hn is not None:                           # 묶음 ㉞ A1
        for sheet in sorted(set(ho) | set(hn)):
            added = sorted(set(hn.get(sheet, [])) - set(ho.get(sheet, [])))
            if added:
                out.append(f"시트 '{sheet}'에 값이 있는 숨긴 행이 새로 생겼습니다: {added}")
    co, cn = old.get("condition_counts"), new.get("condition_counts")
    if co is not None and cn is not None:                           # 묶음 ㉞ A2
        added = sorted(set(cn) - set(co))
        if added:
            out.append(f"새로운 적용 조건 값이 생겼습니다: {added}")
    po, pn = old["per_question"], new["per_question"]
    if set(pn) - set(po):
        out.append(f"새로 생긴 질문: {sorted(set(pn) - set(po))}")
    if set(po) - set(pn):
        out.append(f"사라진 질문: {sorted(set(po) - set(pn))}")
    for qid in sorted(set(po) & set(pn)):
        if po[qid]["digest"] != pn[qid]["digest"]:
            a, b = po[qid], pn[qid]
            detail = f"필수 {a['required']} -> {b['required']}, 보조 {a['optional']} -> {b['optional']}" \
                if (a["required"], a["optional"]) != (b["required"], b["optional"]) else "개수는 같고 내용(유형·질문·근거 위치)이 다릅니다"
            out.append(f"질문 {qid} 내용이 바뀌었습니다: {detail}")
    return out


def _paths(out_dir: str):
    return os.path.join(out_dir, JSON_NAME), os.path.join(out_dir, SNAP_NAME)


def _read_snapshot(out_dir: str):
    p = _paths(out_dir)[1]
    if not os.path.exists(p):
        return None
    with open(p, encoding="utf-8") as f:
        return json.load(f)


def _write_both(out_dir: str, questions: list, snap: dict):
    os.makedirs(out_dir, exist_ok=True)
    jp, sp = _paths(out_dir)
    tmp = []
    for path, data in ((jp, _json_bytes(questions)), (sp, json.dumps(snap, ensure_ascii=False, indent=1).encode("utf-8") + b"\n")):
        with open(path + ".tmp", "wb") as f:
            f.write(data)
        tmp.append((path + ".tmp", path))
    for a, b in tmp:                                                # 검사가 모두 끝난 뒤에만 덮어쓴다
        os.replace(a, b)


def convert(xlsx_path: str, out_dir: str, accept_change: dict | None = None, today: str | None = None, profile: str | None = None) -> dict:
    """정답지를 읽어 JSON과 스냅샷을 쓴다. accept_change={"by": 이름, "reason": 사유}는 확인된 정답지가 달라졌을 때만 쓴다.
    profile은 양식이 다른 정답지(평가셋 v0.5)를 읽는 설정 파일 경로다(묶음 ㉒ V1). 주지 않으면 기존 동작 그대로다."""
    if profile is None:
        questions, report = read_gold_xlsx(xlsx_path)
        report["hidden_rows"] = hidden_rows_report(xlsx_path, [SHEET_Q, SHEET_E], FIRST_ROW)
        report["conditions"] = condition_report(questions)
        new = build_snapshot(xlsx_path, questions, unused_columns=report["unused_columns"],
                             hidden_rows=report["hidden_rows"], condition_counts=report["conditions"]["counts"])
    else:
        pname = profile_name(profile)
        if pname not in _PROFILE_NAMES:
            raise XlsxGoldError(f"{profile}: 모르는 읽기 설정 이름입니다(profile {pname!r}, 아는 이름 {list(_PROFILE_NAMES)})")
        if pname == "b-v0.1":                                     # 묶음 72 D2: 이 이름일 때만 새 함수(v0.5·v0.2 경로는 그대로)
            reader, loader = read_gold_profile_fmtb01, load_profile_fmtb01
        elif pname == "c-v0.1":                                  # 묶음 82 S1: 서비스 C(근거 위치 '시트!N행')
            reader, loader = read_gold_profile_fmtc01, load_profile_fmtc01
        elif pname == "e-v0.2":                                   # 묶음 96 F4: 서비스 E(대체 묶음 값 사용, 30문항 전부)
            reader, loader = read_gold_profile_fmte02, load_profile_fmte02
        elif pname == "d-v0.1":                                # 묶음 97 W7: 서비스 D(서비스 C 양식, 쟁점 없음, 30문항 전부)
            reader, loader = read_gold_profile_fmtd01, load_profile_fmtd01
        else:
            reader = read_gold_profile_v02 if pname == "v0.2" else read_gold_profile
            loader = load_profile_v02 if pname == "v0.2" else load_profile
        questions, report = reader(xlsx_path, profile)
        prof = loader(profile)
        report["hidden_rows"] = hidden_rows_report(xlsx_path, list(prof["sheets"].values()), prof["first_row"])
        report["conditions"] = condition_report(questions)
        new = build_snapshot(xlsx_path, questions, report["profile"], report["unused_columns"],
                             report["hidden_rows"], report["conditions"]["counts"])
    old = _read_snapshot(out_dir)
    result = {"questions": new["questions"], "evidence": new["evidence"], "by_type": new["by_type"], "report": report}
    if old is not None and old["status"] == CONFIRMED:
        diffs = diff_snapshots(old, new)
        if not diffs:
            if not os.path.exists(_paths(out_dir)[0]) or _file_sha(_paths(out_dir)[0]) != old["json_sha256"]:
                raise GoldChanged("확인된 정답지와 같지만 정답표 JSON이 없거나 확인 때와 다릅니다. JSON을 지우고 다시 변환하세요")
            return dict(result, status=CONFIRMED, changed=[])
        if not accept_change or not str(accept_change.get("by", "")).strip() or not str(accept_change.get("reason", "")).strip():
            raise GoldChanged("확인된 정답지와 달라졌습니다. 사람이 확인하고 이름과 사유를 적어 받아들이기 전에는 쓰지 않습니다:\n- " + "\n- ".join(diffs))
        new.update(status=CONFIRMED, confirmed_by=str(accept_change["by"]).strip(), reason="변경 승인: " + str(accept_change["reason"]).strip(),
                   confirmed_date=today or datetime.date.today().isoformat())
        _write_both(out_dir, questions, new)
        return dict(result, status=CONFIRMED, changed=diffs)
    _write_both(out_dir, questions, new)                            # 처음이거나 아직 확인 전: 미확인으로 쓴다
    return dict(result, status=UNCONFIRMED, changed=[])


def confirm(out_dir: str, by: str, reason: str, today: str | None = None) -> dict:
    """사람이 정답표를 확인했다고 기록한다. JSON이 스냅샷과 같을 때만 된다."""
    if not str(by).strip() or not str(reason).strip():
        raise XlsxGoldError("확인한 사람과 사유가 필요합니다")
    snap = _read_snapshot(out_dir)
    if snap is None:
        raise GoldNotConfirmed("확인할 스냅샷이 없습니다. 먼저 변환하세요")
    jp = _paths(out_dir)[0]
    if not os.path.exists(jp) or _file_sha(jp) != snap["json_sha256"]:
        raise GoldNotConfirmed("정답표 JSON이 스냅샷과 다릅니다. 다시 변환하세요")
    if snap["status"] == CONFIRMED:
        raise XlsxGoldError("이미 확인된 스냅샷입니다")
    snap.update(status=CONFIRMED, confirmed_by=str(by).strip(), reason=str(reason).strip(), confirmed_date=today or datetime.date.today().isoformat())
    with open(_paths(out_dir)[1], "w", encoding="utf-8") as f:
        json.dump(snap, f, ensure_ascii=False, indent=1)
        f.write("\n")
    return snap


def load_confirmed(out_dir: str) -> list:
    """판정과 실행기가 정답표를 읽는 입구(N9). 확인됨이고 JSON이 확인 때와 같을 때만 질문 목록을 돌려준다."""
    snap = _read_snapshot(out_dir)
    if snap is None:
        raise GoldNotConfirmed(f"{out_dir}: 스냅샷이 없습니다. 먼저 변환하고 사람이 확인해야 합니다")
    if snap["status"] != CONFIRMED:
        raise GoldNotConfirmed(f"{out_dir}: 정답표가 아직 '{snap['status']}' 상태입니다. 사람이 확인하기 전에는 판정에 쓰지 않습니다")
    jp = _paths(out_dir)[0]
    if not os.path.exists(jp) or _file_sha(jp) != snap["json_sha256"]:
        raise GoldNotConfirmed(f"{out_dir}: 정답표 JSON이 확인 때와 다릅니다(손으로 고쳤거나 바뀜). 다시 변환하고 확인하세요")
    return scoring.load_gold(jp)


# ---------- 명령줄 ----------
def _main(argv):
    import argparse

    ap = argparse.ArgumentParser(description="정답지 엑셀 -> 정답표 JSON, 변경 감지")
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("convert")
    c.add_argument("xlsx")
    c.add_argument("out_dir")
    c.add_argument("--accept-change", action="store_true", help="확인된 정답지가 달라졌을 때 사람이 받아들임(--by, --reason 필요)")
    c.add_argument("--by", default="")
    c.add_argument("--reason", default="")
    c.add_argument("--format", default=None, help="양식이 다른 정답지(평가셋 v0.5)를 읽는 설정 파일(예: gold_format_v05.yaml). 없으면 기존 양식")
    k = sub.add_parser("confirm")
    k.add_argument("out_dir")
    k.add_argument("--by", required=True)
    k.add_argument("--reason", required=True)
    a = ap.parse_args(argv)
    try:
        if a.cmd == "convert":
            r = convert(a.xlsx, a.out_dir, {"by": a.by, "reason": a.reason} if a.accept_change else None, profile=a.format)
            print(f"질문 {r['questions']}개, 근거 {r['evidence']}행, 유형별 {r['by_type']}, 상태 {r['status']}")
            print("사용 안 한 열:", r["report"]["unused_columns"])
            hidden = {k: v for k, v in r["report"].get("hidden_rows", {}).items() if v}
            if hidden:                                               # 묶음 ㉞ A1: 막지 않고 알린다(삭제로 볼지는 사람이 정함)
                print("[경고] 값이 있는 숨긴 행이 읽혔습니다(읽는 범위는 그대로):", hidden)
            cond = r["report"].get("conditions") or {}
            if cond.get("unstated"):                                 # 묶음 ㉞ A2
                print("[경고] 질문에 명시되지 않았고 '항상'도 아닌 적용 조건 값(질문 id별):", cond["unstated"])
            p = r["report"].get("profile")
            if p:                                                    # 묶음 ㉒ V22·V23: 건수와 상태만 적는다(질문 글 없음)
                print(f"읽기 설정 {p['name']}: 보류 전 질문 {p['n_questions_before_hold']}개, 근거 {p['n_evidence_before_hold']}행, "
                      f"단위(필수 {p['units_before_hold']['required']}, 보조 {p['units_before_hold']['optional']}), 단위 모양 {p['unit_shapes']}")
                print("문서 상태:", {k: v["status"] for k, v in p["documents"].items()})
                print("보류 질문:", [(h["id"], h["documents"]) for h in p["held_questions"]] or "없음")
                print("기획팀 쪽 상태:", p["source_state"])
                if p.get("not_available"):                           # 묶음 ㉖ W3·W6: 열이 없어 정하지 않은 값
                    print("열 없음:", p["not_available"])
                if p.get("question_subset"):                         # 묶음 72 D1c: 고른 결과와 뺀 질문 ID(질문 글 없음)
                    qs_ = p["question_subset"]
                    print(f"질문 고르기 {qs_['subset']}: 뺀 질문 {qs_['left_out']}, 빼기 전 개수 {qs_['before_subset']}")
            if r["changed"]:
                print("승인한 변경:\n- " + "\n- ".join(r["changed"]))
        else:
            s = confirm(a.out_dir, a.by, a.reason)
            print(f"확인됨: {s['confirmed_by']} {s['confirmed_date']}")
    except XlsxGoldError as err:
        print(f"멈춤: {err}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(_main(sys.argv[1:]))
