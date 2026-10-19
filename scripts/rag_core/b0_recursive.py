"""일반 리커시브 기준선(B0-RC) 청크 만들기(묶음 90, 기준은 CHANGES.md 묶음 90 R1~R6). 기준선 3개(B0-RC / B0-S / B0-R) 중 하나다.
(처음 서비스 A·서비스 B·서비스 C에서는 결과를 본 뒤 잰 [사후] 참고였고, 서비스 E·서비스 D에서는 결과 전에 등록해 다섯 비교에 썼다.)

엑셀의 모든 시트·모든 행을 프로필 없이 읽어, 행마다 빈 칸이 아닌 값을 " | "로 잇고(머리글을 칸에 붙이지 않음) 줄바꿈으로 이은 뒤
LangChain RecursiveCharacterTextSplitter(500자, 겹침 50자)로 시트마다 자른다. 평가기는 B0-R/B0-S 두 칸만 알므로
청크 id는 B0-S 칸(`B0-S-…`)을 쓴다(내용은 일반 리커시브, 묶음 90 R3).
실행: python b0_recursive.py --files <엑셀 ...> --out <저장소 밖 폴더>
"""
from __future__ import annotations

import argparse
import json
import os
import sys

import openpyxl

from b0_chunks import _schema_version, clean

CHUNK_SIZE, CHUNK_OVERLAP = 500, 50
SEPARATORS = ["\n\n", "\n", " ", ""]
PREFIX = "B0-S"                                    # R3: 평가기의 B0-S 칸을 빌림
OUT_NAME = "b0_rc_chunks.jsonl"


class RecursiveError(RuntimeError):
    """청크 위치를 정할 수 없는 등 기준(R1)을 지킬 수 없을 때."""


def sheet_lines(ws) -> list:
    """(행, 글) 목록. 모든 행을 읽고 빈 칸이 아닌 값을 " | "로 잇는다. 머리글을 붙이지 않는다. 모든 칸이 빈 행은 건너뛴다."""
    out = []
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, values_only=True), 1):   # 1행부터 세야 행 번호가 맞다
        vals = [v for v in (clean(x) for x in row) if v is not None]
        if vals:
            out.append((r, " | ".join(str(v) for v in vals)))
    return out


def _splitter():
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    return RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP, separators=SEPARATORS,
                                          length_function=len, add_start_index=True)


def split_sheet(lines: list, splitter=None) -> list:
    """한 시트의 (행, 글) 목록을 이어 자르고 [(청크 글, (첫 행, 끝 행))]를 돌려준다. 행 범위는 청크 글이 걸친 행."""
    if not lines:
        return []
    spans, parts, pos = [], [], 0
    for row, text in lines:
        spans.append((row, pos, pos + len(text)))
        parts.append(text)
        pos += len(text) + 1                                   # 줄바꿈 한 글자
    full = "\n".join(parts)
    out = []
    for d in (splitter or _splitter()).create_documents([full]):
        start = d.metadata.get("start_index", -1)
        if start is None or start < 0 or full[start:start + len(d.page_content)] != d.page_content:
            raise RecursiveError("분할기가 청크 시작 위치를 찾지 못했습니다(행 범위를 정할 수 없어 멈춤)")
        end = start + len(d.page_content)
        rows = [r for r, s, e in spans if s < end and e > start]
        if not rows:
            raise RecursiveError("청크가 걸친 행을 찾지 못했습니다")
        out.append((d.page_content, (min(rows), max(rows))))
    return out


def build_b0_recursive(files: list, out_dir: str) -> dict:
    """파일마다 모든 시트를 읽어 청크 파일 하나를 쓴다. source.file에는 파일 이름만 쓴다(B0·B1과 같은 표기)."""
    if not out_dir:
        raise RecursiveError("out_dir이 필요합니다")
    names = [os.path.basename(f) for f in files]
    if len(set(names)) != len(names):
        raise RecursiveError(f"같은 파일 이름이 둘 이상입니다: {names}")
    splitter, chunks, report = _splitter(), [], []
    for path in files:
        if not os.path.isfile(path):
            raise RecursiveError(f"파일이 없습니다: {path}")
        wb = openpyxl.load_workbook(path, data_only=True)
        for sheet in wb.sheetnames:                           # R1 (가): 숨김 포함 모든 시트
            ws = wb[sheet]
            lines = sheet_lines(ws)
            parts = split_sheet(lines, splitter)
            for text, rows in parts:
                chunks.append((os.path.basename(path), sheet, rows, text))
            report.append({"file": os.path.basename(path), "sheet": sheet, "hidden_sheet": ws.sheet_state != "visible",
                           "rows_with_text": len(lines), "chunks": len(parts)})
    version = _schema_version()
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, OUT_NAME)
    if os.path.exists(path):
        raise RecursiveError(f"이미 있는 파일입니다(덮어쓰지 않음): {path}")
    with open(path, "x", encoding="utf-8", newline="\n") as f:
        for i, (file_label, sheet, rows, text) in enumerate(chunks, 1):
            f.write(json.dumps({"schema_version": version, "chunk_id": f"{PREFIX}-{i:06d}", "text": text,
                                "source": {"file": file_label, "sheet": sheet, "rows": {"start": rows[0], "end": rows[1]}},
                                "meta": {}}, ensure_ascii=False) + "\n")
    lens = [len(c[3]) for c in chunks]
    return {"path": path, "n_chunks": len(chunks), "mean_chars": sum(lens) / len(lens) if lens else 0, "max_chars": max(lens, default=0),
            "chunk_size": CHUNK_SIZE, "chunk_overlap": CHUNK_OVERLAP, "separators": SEPARATORS, "sheets": report}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--files", nargs="+", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    r = build_b0_recursive(a.files, a.out)
    with open(os.path.join(a.out, "b0_rc_report.json"), "x", encoding="utf-8") as f:
        json.dump(r, f, ensure_ascii=False, indent=1)
    print("청크", r["n_chunks"], "평균 글자", round(r["mean_chars"], 1), "최대", r["max_chars"], "시트", len(r["sheets"]))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except RecursiveError as e:
        print(f"[멈춤] {e}", file=sys.stderr)
        sys.exit(2)
