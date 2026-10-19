"""엑셀 원본에서 B0 청크(B0-R, B0-S)를 만들어 chunks.jsonl로 쓴다.

[쉽게 말하면]
  기준선 카드를 만드는 곳이다. 엑셀을 한 행씩 읽어 "열이름: 값 | 열이름: 값" 한 줄을 만든 뒤 두 가지로 쓴다.
    - B0-R : 그 한 줄이 카드 하나(행 하나 = 카드 하나). 카드가 짧아 같은 예산에 많이 담긴다.
    - B0-S : 그 줄들을 줄바꿈으로 이어 500자씩(겹침 50자) 자른 카드. 행 중간에서도 끊긴다.
  어느 시트를 읽을지, 머리글이 몇 행인지, 데이터가 몇 행부터인지는 si-prep 프로필(yaml)을 읽어 따른다.
  header_mode="none"이면 열 이름을 붙이지 않고 1행부터 읽는다(머리글 행도 한 줄로 들어감).
  병합 셀 복원, 분류 경로, 문서 간 연결은 하지 않는다(그건 si-prep, 즉 B1이 하는 일).

[이 파일에서 먼저 볼 것]
  build_b0      : 입구. 프로젝트 대응표(파일 ↔ 프로필)를 받아 두 청크 파일을 쓴다
  row_text      : 한 행 → "열이름: 값 | …" 한 줄
  split_windows : 줄들을 500자·겹침 50자로 자름(B0-S)

기준은 CHANGES.md의 '묶음 ③'(G1~G10). B0는 전처리를 거치지 않은 기준선이라 병합 복원도 분류 경로도 없다.
si-prep의 프로필 YAML은 읽기만 한다. si-prep 코드에서는 버전 문자열(`SCHEMA_VERSION`)만 불러온다(INTERFACE.md, 묶음 ⑬).
"""
from __future__ import annotations

import json
import os
import re

import openpyxl
import yaml

HEADER_MODES = ("profile", "none")
WINDOW = 500
OVERLAP = 50


class B0Error(RuntimeError):
    """입력이 기준(G8)에 맞지 않을 때."""


def _schema_version() -> str:
    """팀 스키마의 버전 문자열을 그대로 쓴다. 우리가 따로 적으면 스키마가 오를 때 어긋난다(묶음 ⑬)."""
    try:
        from si_prep.schema import SCHEMA_VERSION
    except ImportError as e:
        raise B0Error("si_prep을 불러올 수 없습니다. 팀 저장소를 pip install -e로 설치하세요(INTERFACE.md).") from e
    return SCHEMA_VERSION


def clean(v):
    """셀 값을 문자열로. 빈 값은 None. 1.0은 '1'로(si-prep reader.clean과 같은 규칙, G2)."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


# ---------- 프로필 읽기(묶음 ㉔) ----------
# 팀 설계(ARCHITECTURE §5.3)의 키 중 B0가 읽는 것은 "읽기 범위" 키뿐이다. 전처리 판단 키를 읽으면 B0가 B1이 된다.
RANGE_KEYS = {"id", "approved", "sheets", "header_rows", "data_start_row", "stop_markers", "sheet_overrides"}
JUDGMENT_KEYS = {"columns", "hierarchy", "item_boundary", "sheet_as_root", "root_separators", "overview_sheets",
                 "version", "id_pattern", "table_gap_rows", "hidden_rows",
                 "tables"}                                     # 묶음 68(62 S1): 한 시트 여러 표(스키마 1.4)는 B0가 읽지 않고 보고만 한다
OVERRIDE_KEYS = {"sheets", "header_rows", "data_start_row", "stop_markers"}


def _is_new_format(profile: dict) -> bool:
    """새 형식 표지: sheets.include/exclude, stop_markers, sheet_overrides 중 하나. 없으면 옛 방식(모르는 키를 넘긴다)."""
    s = profile.get("sheets")
    if isinstance(s, list):                                       # 팀이 실제로 보낸 형식: [{name, read, reason}]
        return True
    if isinstance(s, dict) and s and set(s) <= {"include", "exclude"}:
        return True
    return "stop_markers" in profile or "sheet_overrides" in profile


def _matches(name: str, pattern: str) -> bool:
    """시트 이름 매칭: 이름 그대로(끝 공백 포함), `*`는 여러 글자, `?`는 한 글자. 대소문자 구분(§4)."""
    rx = "".join(".*" if c == "*" else "." if c == "?" else re.escape(c) for c in pattern)
    return re.fullmatch(rx, name, flags=re.DOTALL) is not None


def _as_list(v, label: str, key: str) -> list:
    if v is None:
        return []
    if not isinstance(v, list) or not all(isinstance(x, str) for x in v):
        raise B0Error(f"{label}: {key}는 글자 목록이어야 합니다: {v!r}")
    return v


def parse_profile(profile: dict, label: str) -> dict:
    """프로필을 한 모양으로 바꾼다. 새 형식(§5.3)과 옛 형식을 모두 읽는다. 둘을 섞으면 멈춘다."""
    profile = profile or {}
    legacy_markers = [k for k in ("include_sheets", "skip_sheets") if k in profile]
    legacy_map = isinstance(profile.get("sheets"), dict) and any(k not in ("include", "exclude") for k in profile["sheets"])
    if _is_new_format(profile):
        if legacy_markers or legacy_map:
            raise B0Error(f"{label}: 새 형식(sheets.include/exclude, stop_markers, sheet_overrides)과 옛 형식"
                          f"({legacy_markers or '시트별 sheets'})을 한 프로필에 섞을 수 없습니다")
        unknown = sorted(set(profile) - RANGE_KEYS - JUDGMENT_KEYS)
        if unknown:
            raise B0Error(f"{label}: 모르는 프로필 키입니다: {unknown}")
        sheets = profile.get("sheets") or {}
        include, exclude, listed = [], [], []
        if isinstance(sheets, list):                              # L1: [{name, read?, reason?}], 기본 읽음
            for i, e in enumerate(sheets):
                if not isinstance(e, dict) or not isinstance(e.get("name"), str) or set(e) - {"name", "read", "reason"}:
                    raise B0Error(f"{label}: sheets[{i}]는 name(글자)과 선택 칸 read, reason만 가질 수 있습니다: {e!r}")
                if "read" in e and not isinstance(e["read"], bool):
                    raise B0Error(f"{label}: sheets[{i}].read는 true/false여야 합니다: {e['read']!r}")
                listed.append(e["name"])
                if e.get("read") is False:
                    exclude.append(e["name"])
        else:
            include = _as_list(sheets.get("include"), label, "sheets.include")
            exclude = _as_list(sheets.get("exclude"), label, "sheets.exclude")
        overrides, ignored = [], set(profile) & JUDGMENT_KEYS
        for i, ov in enumerate(profile.get("sheet_overrides") or []):
            if not isinstance(ov, dict):
                raise B0Error(f"{label}: sheet_overrides[{i}]가 항목(키: 값)이 아닙니다: {ov!r}")
            unknown = set(ov) - OVERRIDE_KEYS - JUDGMENT_KEYS         # L2: 판단 키는 읽지 않고 보고, 모르는 키는 멈춤
            if unknown:
                raise B0Error(f"{label}: sheet_overrides[{i}]에 허용되지 않는 키가 있습니다: {sorted(unknown)}")
            ignored |= {f"sheet_overrides.{k}" for k in set(ov) & JUDGMENT_KEYS}
            pats = _as_list(ov.get("sheets"), label, f"sheet_overrides[{i}].sheets")
            if not pats:
                raise B0Error(f"{label}: sheet_overrides[{i}]에 sheets가 필요합니다")
            overrides.append((pats, {k: ov[k] for k in ov if k in OVERRIDE_KEYS and k != "sheets"}))
        return {"format": "v53", "include": include, "exclude": exclude, "listed": listed,
                "defaults": {k: profile.get(k) for k in ("header_rows", "data_start_row", "stop_markers")},
                "overrides": overrides, "legacy_sheet_map": {},
                "ignored_keys": sorted(ignored)}
    return {"format": "legacy", "include": list(profile.get("include_sheets") or []),
            "exclude": list(profile.get("skip_sheets") or []), "exact_include": True,
            "defaults": {k: profile.get(k) for k in ("header_rows", "data_start_row", "stop_markers")},
            "overrides": [], "legacy_sheet_map": dict(profile.get("sheets") or {}), "ignored_keys": []}


def _target_sheets(wb, parsed: dict, label: str) -> list[str]:
    """읽을 시트와 보고. 새 형식: include가 있으면 맞는 시트만, exclude에 맞는 시트는 뺀다(문서 순서)."""
    names = wb.sheetnames
    missing_listed = [n for n in parsed.get("listed") or [] if n not in names]
    if missing_listed:
        raise B0Error(f"{label}: 프로필 sheets 목록에 있는 이름이 문서에 없습니다: {missing_listed}")
    if parsed.get("exact_include"):                               # 옛 방식: 이름 그대로, 적은 순서
        if parsed["include"]:
            missing = [s for s in parsed["include"] if s not in names]
            if missing:
                raise B0Error(f"{label}: 프로필의 include_sheets에 있는 시트가 문서에 없습니다: {missing}")
            return list(parsed["include"])
        skip = set(parsed["exclude"])
        return [s for s in names if s not in skip]
    chosen = names
    if parsed["include"]:
        dead = [p for p in parsed["include"] if not any(_matches(n, p) for n in names)]
        if dead:
            raise B0Error(f"{label}: sheets.include에 맞는 시트가 문서에 없습니다: {dead}")
        chosen = [n for n in names if any(_matches(n, p) for p in parsed["include"])]
    return [n for n in chosen if not any(_matches(n, p) for p in parsed["exclude"])]


def _unmatched_exclude(wb, parsed: dict) -> list[str]:
    return [p for p in parsed["exclude"] if not any(_matches(n, p) for n in wb.sheetnames)]


def _sheet_config(parsed: dict, sheet: str, label: str) -> dict:
    """한 시트의 header_rows, data_start_row, stop_markers. 덮어쓰기는 적은 순서대로 적용한다."""
    cfg = dict(parsed["defaults"])
    if parsed["format"] == "legacy":
        cfg.update({k: v for k, v in (parsed["legacy_sheet_map"].get(sheet) or {}).items() if k in cfg})
    else:
        for pats, ov in parsed["overrides"]:
            if any(_matches(sheet, p) for p in pats):
                cfg.update(ov)
    hr, start = cfg.get("header_rows"), cfg.get("data_start_row")
    if hr and start and start <= max(hr):
        raise B0Error(f"{label} / {sheet}: data_start_row({start})는 max(header_rows)({max(hr)})보다 커야 합니다")
    cfg["stop_markers"] = _as_list(cfg.get("stop_markers"), f"{label} / {sheet}", "stop_markers")
    return cfg


def _first_value(ws, row: int):
    for col in range(1, ws.max_column + 1):
        v = clean(ws.cell(row, col).value)
        if v is not None:
            return v
    return None


def _headers(ws, header_rows: list[int]) -> dict[int, str]:
    out = {}
    for col in range(1, ws.max_column + 1):
        parts = [v for r in header_rows if (v := clean(ws.cell(r, col).value))]
        if parts:
            out[col] = " / ".join(parts)
    return out


FORMULA_CELL_LIMIT = 20            # 보고에 좌표를 남기는 최대 개수


def find_uncached_formulas(ws_formula, ws_value, first_row: int, last_row: int) -> list:
    """수식 칸인데 저장된 계산값이 없는 칸의 (행, 좌표) 목록(Q1). 값은 만들어 채우지 않는다.

    한계: 수식의 결과가 빈 문자열인 경우와 계산값이 저장되지 않은 경우를 구분하지 못해 둘 다 센다.
    """
    out = []
    for row in ws_formula.iter_rows(min_row=first_row, max_row=last_row):
        for c in row:
            if c.data_type == "f" and ws_value[c.coordinate].value is None:
                out.append((c.row, c.coordinate))
    return out


def hidden_columns(ws) -> set:
    """숨김 열 번호(1부터). openpyxl은 속성이 같은 이웃 열을 한 항목(min~max)으로 묶어 두므로 펼친다.
    시트 최대 열까지만 센다(오른쪽 끝까지 숨긴 범위가 수천 열로 세어지지 않게, 묶음 70)."""
    cols = set()
    for key, d in ws.column_dimensions.items():
        if d.hidden:
            lo = d.min or openpyxl.utils.column_index_from_string(key)
            cols.update(range(lo, min(d.max or lo, ws.max_column) + 1))
    return cols


def row_text(ws, row: int, headers: dict[int, str] | None) -> str | None:
    """한 행의 텍스트(G1). 모든 칸이 비어 있으면 None."""
    parts = []
    for col in range(1, ws.max_column + 1):
        v = clean(ws.cell(row, col).value)
        if v is None:
            continue
        h = headers.get(col) if headers else None
        parts.append(f"{h}: {v}" if h else v)
    return " | ".join(parts) if parts else None


def split_windows(lines: list[tuple[int, str]], window: int = WINDOW, overlap: int = OVERLAP):
    """(행, 텍스트) 목록을 줄바꿈으로 이어 window자·overlap자 겹침으로 자른다(G5). (text, (첫 행, 끝 행))를 돌려준다."""
    if not lines:
        return []
    if not 0 <= overlap < window:
        raise B0Error("overlap은 0 이상 window 미만이어야 합니다")
    full, spans, pos = [], [], 0
    for row, text in lines:
        spans.append((row, pos, pos + len(text)))
        full.append(text)
        pos += len(text) + 1                       # 줄바꿈 한 글자
    full = "\n".join(full)
    out, start, step = [], 0, window - overlap
    while True:
        end = min(start + window, len(full))
        rows = [r for r, s, e in spans if s < end and e > start]
        if not rows:
            raise B0Error("청크가 걸친 행을 찾지 못했습니다(내부 오류)")
        out.append((full[start:end], (min(rows), max(rows))))
        if end == len(full):
            return out
        start += step


def _read_project(project_yaml: str, base_dir: str) -> list[tuple[str, str, dict]]:
    with open(project_yaml, encoding="utf-8") as f:
        proj = yaml.safe_load(f)
    docs = (proj or {}).get("documents")
    if not docs:
        raise B0Error(f"{project_yaml}: documents가 비어 있습니다")
    result = []
    for d in docs:
        file_label, profile_path = d.get("file"), d.get("profile")
        if not file_label or not profile_path:
            raise B0Error(f"{project_yaml}: documents 항목에 file과 profile이 모두 있어야 합니다: {d}")
        with open(os.path.join(base_dir, profile_path), encoding="utf-8") as f:
            # source.file에는 파일 이름만 쓴다(정답표가 한 가지 표기만 쓰므로 B0와 B1이 같아야 한다). 열 때는 전체 경로를 쓴다.
            result.append((re.split(r"[\\/]", file_label)[-1], os.path.join(base_dir, file_label), yaml.safe_load(f) or {}))
    names = [r[0] for r in result]
    dup = sorted({n for n in names if names.count(n) > 1})
    if dup:
        raise B0Error(f"{project_yaml}: 같은 파일 이름이 둘 이상입니다(정답표가 구분할 수 없음): {dup}")
    return result


def build_b0(project_yaml: str, base_dir: str, out_dir: str, header_mode: str, meta_prefix: bool = False) -> dict:
    """B0-R과 B0-S 청크 파일을 out_dir에 쓰고 보고를 돌려준다.

    base_dir: 프로젝트 YAML의 상대 경로(파일·프로필)를 푸는 기준 폴더. source.file에는 파일 이름만 쓴다(G4).
    """
    if header_mode not in HEADER_MODES:
        raise B0Error(f"header_mode는 {HEADER_MODES} 중 하나여야 합니다: {header_mode!r}")
    if not out_dir:
        raise B0Error("out_dir이 필요합니다(기본값 없음, G9)")

    chunks_r, chunks_s, sheets_report, profiles_report = [], [], [], []
    for file_label, path, profile in _read_project(project_yaml, base_dir):
        if not os.path.isfile(path):
            raise B0Error(f"파일이 없습니다: {path}")
        parsed = parse_profile(profile, file_label)
        wb = openpyxl.load_workbook(path, data_only=True)
        wbf = openpyxl.load_workbook(path, data_only=False)          # Q1: 수식 칸을 찾기 위해 한 번 더 읽는다
        profiles_report.append({"file": file_label, "format": parsed["format"], "ignored_keys": parsed["ignored_keys"],
                                "unmatched_exclude": _unmatched_exclude(wb, parsed),
                                "unlisted_sheets": [n for n in wb.sheetnames if parsed.get("listed") and n not in parsed["listed"]]})
        for sheet in _target_sheets(wb, parsed, file_label):
            ws = wb[sheet]
            cfg = _sheet_config(parsed, sheet, file_label)
            if header_mode == "profile":
                header_rows, start = cfg["header_rows"], cfg["data_start_row"]
                if not header_rows or not start:
                    raise B0Error(f"{file_label} / {sheet}: 프로필에 header_rows와 data_start_row가 필요합니다")
                headers = _headers(ws, list(header_rows))
            else:
                headers, start = None, 1
            uncached = find_uncached_formulas(wbf[sheet], ws, min(header_rows) if header_mode == "profile" else 1, ws.max_row)
            uncached_rows = {r for r, _ in uncached}
            lines, skipped, lost, stopped_at = [], 0, 0, None
            for row in range(start, ws.max_row + 1):
                if cfg["stop_markers"]:
                    first = _first_value(ws, row)
                    if first is not None and any(first.startswith(m) for m in cfg["stop_markers"]):
                        stopped_at = row                  # 표 끝 표시: 이 행부터 데이터 끝(§5.3)
                        break
                text = row_text(ws, row, headers)
                if text is None:
                    skipped += 1
                    if row in uncached_rows:
                        lost += 1                       # Q2: 값이 전부 계산값 없는 수식이라서 빈 행으로 건너뛴 행
                else:
                    lines.append((row, text))
            windows = split_windows(lines)
            stem = os.path.splitext(file_label)[0]
            for row, text in lines:
                if meta_prefix:                               # 묶음 ㉘ X2(C1): 파일·시트 이름 머리말 한 줄. 꺼짐이면 기존 청크와 같다
                    text = f"[파일: {stem} | 시트: {sheet}]" + "\n" + text
                chunks_r.append((file_label, sheet, (row, row), text))
            for text, rows in windows:
                chunks_s.append((file_label, sheet, rows, text))
            hidden_n = sum(1 for r in range(start, ws.max_row + 1)
                           if ws.row_dimensions[r].hidden and any(c.value not in (None, "") for c in ws[r]))      # 묶음 ㉞ A3: 보고만 한다
            hid = sorted(hidden_columns(ws))
            hidden_cells = [(r, c) for r, _ in lines for c in hid if clean(ws.cell(r, c).value) is not None]   # 묶음 70: 보고만 한다(이슈 34)
            sheets_report.append({"file": file_label, "sheet": sheet, "data_rows": len(lines),
                                  "hidden_sheet": ws.sheet_state != "visible", "hidden_rows": hidden_n,
                                  "hidden_cols": len({c for _, c in hidden_cells}), "hidden_col_cells": len(hidden_cells),
                                  "skipped_empty_rows": skipped, "b0_r": len(lines), "b0_s": len(windows),
                                  "stopped_at_marker_row": stopped_at,
                                  "formula_no_value": len(uncached),
                                  "formula_no_value_cells": [c for _, c in uncached[:FORMULA_CELL_LIMIT]],
                                  "rows_lost_to_formula_no_value": lost})

    version = _schema_version()                       # 파일을 쓰기 전에 확인한다(못 불러오면 아무것도 쓰지 않고 멈춤)
    os.makedirs(out_dir, exist_ok=True)
    paths = {}
    for prefix, name, chunks in (("B0-R", "b0_r_chunks.jsonl", chunks_r), ("B0-S", "b0_s_chunks.jsonl", chunks_s)):
        paths[prefix] = os.path.join(out_dir, name)
        with open(paths[prefix], "w", encoding="utf-8", newline="\n") as f:
            for i, (file_label, sheet, rows, text) in enumerate(chunks, 1):
                f.write(json.dumps({
                    "schema_version": version, "chunk_id": f"{prefix}-{i:06d}", "text": text,
                    "source": {"file": file_label, "sheet": sheet, "rows": {"start": rows[0], "end": rows[1]}}, "meta": {},
                }, ensure_ascii=False) + "\n")
    return {"header_mode": header_mode, "meta_prefix": bool(meta_prefix), "n_b0_r": len(chunks_r), "n_b0_s": len(chunks_s),
            "n_formula_no_value": sum(x["formula_no_value"] for x in sheets_report),
            "n_rows_lost_to_formula_no_value": sum(x["rows_lost_to_formula_no_value"] for x in sheets_report),
            "n_hidden_cols": sum(x["hidden_cols"] for x in sheets_report), "n_hidden_col_cells": sum(x["hidden_col_cells"] for x in sheets_report),
            "paths": paths, "sheets": sheets_report, "profiles": profiles_report}
